"""Safe text extraction for administrator-supplied Atlas knowledge files."""

from __future__ import annotations

import io
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import unquote

from docx import Document


ATLAS_KNOWLEDGE_MAX_FILE_BYTES = 8 * 1024 * 1024
ATLAS_KNOWLEDGE_MAX_TEXT_CHARS = 250_000
ATLAS_KNOWLEDGE_EXTENSIONS = frozenset({".txt", ".md", ".docx", ".pdf"})
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class AtlasKnowledgeFileError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = str(code)


def _clean_text(value: str) -> str:
    text = _CONTROL_RE.sub("", str(value or "")).replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    clean = "\n".join(lines).strip()
    if len(clean) < 20:
        raise AtlasKnowledgeFileError(
            "atlas_file_text_empty",
            "В файле не найден читаемый текст. Для скана используйте распознанную копию.",
        )
    if len(clean) > ATLAS_KNOWLEDGE_MAX_TEXT_CHARS:
        raise AtlasKnowledgeFileError(
            "atlas_file_text_too_large",
            "Материал слишком большой. Разделите его на несколько самостоятельных источников.",
        )
    return clean


def _plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            return _clean_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise AtlasKnowledgeFileError(
        "atlas_file_encoding_unsupported",
        "Не удалось прочитать кодировку файла. Сохраните его как UTF-8.",
    )


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded_size = sum(item.file_size for item in archive.infolist())
            if expanded_size > 40 * 1024 * 1024:
                raise AtlasKnowledgeFileError(
                    "atlas_file_docx_too_large",
                    "Содержимое DOCX слишком большое. Разделите документ на несколько материалов.",
                )
        document = Document(io.BytesIO(data))
    except AtlasKnowledgeFileError:
        raise
    except Exception as exc:
        raise AtlasKnowledgeFileError(
            "atlas_file_docx_invalid",
            "DOCX повреждён или имеет неподдерживаемый формат.",
        ) from exc
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return _clean_text("\n\n".join(blocks))


def _pdf_text(data: bytes) -> str:
    with tempfile.TemporaryDirectory(prefix="atlas-pdf-") as directory:
        source = Path(directory) / "source.pdf"
        source.write_bytes(data)
        try:
            completed = subprocess.run(
                ["pdftotext", "-layout", "-nopgbrk", str(source), "-"],
                capture_output=True,
                check=False,
                timeout=20,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise AtlasKnowledgeFileError(
                "atlas_file_pdf_unavailable",
                "PDF сейчас не удалось обработать. Сохраните материал как DOCX или TXT.",
            ) from exc
    if completed.returncode != 0:
        raise AtlasKnowledgeFileError(
            "atlas_file_pdf_invalid",
            "PDF повреждён, защищён или не содержит доступного текста.",
        )
    return _plain_text(completed.stdout)


def atlas_extract_knowledge_file(filename: str, data: bytes) -> dict[str, str | int]:
    clean_name = unquote(str(filename or "")).replace("\\", "/").rsplit("/", 1)[-1][:240]
    suffix = Path(clean_name).suffix.lower()
    if not clean_name or suffix not in ATLAS_KNOWLEDGE_EXTENSIONS:
        raise AtlasKnowledgeFileError(
            "atlas_file_type_unsupported",
            "Поддерживаются TXT, Markdown, DOCX и PDF.",
        )
    if not data:
        raise AtlasKnowledgeFileError("atlas_file_empty", "Выбранный файл пуст.")
    if len(data) > ATLAS_KNOWLEDGE_MAX_FILE_BYTES:
        raise AtlasKnowledgeFileError(
            "atlas_file_too_large",
            "Файл превышает лимит 8 МБ.",
        )
    if suffix in {".txt", ".md"}:
        text = _plain_text(data)
    elif suffix == ".docx":
        text = _docx_text(data)
    else:
        text = _pdf_text(data)
    return {
        "filename": clean_name,
        "title": Path(clean_name).stem[:180],
        "content": text,
        "size": len(data),
    }


__all__ = [
    "ATLAS_KNOWLEDGE_EXTENSIONS",
    "ATLAS_KNOWLEDGE_MAX_FILE_BYTES",
    "AtlasKnowledgeFileError",
    "atlas_extract_knowledge_file",
]
