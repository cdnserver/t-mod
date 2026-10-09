import asyncio
import base64
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import storage
from aiohttp import FormData, web
from aiohttp.test_utils import TestClient, TestServer
from docx import Document
from modules.atlas_ai import (
    AtlasAIConfig,
    AtlasAIError,
    _answer_text,
    _atlas_answer_is_retrieval_refusal,
    _atlas_corpus_abbreviations,
    _atlas_merge_source_fragments,
    _atlas_pinpoint_labels,
    _atlas_query_variants,
    _atlas_relevant_sources,
    _atlas_task_profile,
    _bounded_dialog_messages,
    _cross_chat_context,
    _chunks,
    _compact_overlay_answer,
    _deterministic_exact_lookup,
    _grounded_refusal_fallback,
    _recent_user_dialog_context,
    _response_delivery_contract,
    atlas_ai_config,
    atlas_answer,
    atlas_answer_stream,
    atlas_embed,
    atlas_ensure_collection,
    atlas_index_source,
    atlas_model_route,
    atlas_parse_text_mode,
    atlas_probe_collection,
    atlas_research_plan,
    atlas_search,
)
from modules.atlas_agents import atlas_agent_catalog, atlas_resolve_agent
from modules.atlas_knowledge import AtlasKnowledgeFileError, atlas_extract_knowledge_file
from modules.atlas_taxonomy import atlas_classify_knowledge
from modules.atlas_forum_sync import AtlasForumScrapeBatch, AtlasForumSnapshot
from modules.atlas_web import register_atlas_web_routes
from modules.consensus_web import create_consensus_web_app
from modules.consensus_web_auth import ConsensusWebPrincipal
from persistence import atlas_repository
from persistence import atlas_forum_attachment_repository
from persistence import atlas_job_repository
from persistence.core import connect


class AtlasRepositoryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.old_data_dir = storage.DATA_DIR
        self.old_database_file = storage.DATABASE_FILE
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(self.temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-test.db"
        storage.init_db()

    def tearDown(self) -> None:
        storage.DATA_DIR = self.old_data_dir
        storage.DATABASE_FILE = self.old_database_file
        self.temp_dir.cleanup()

    def test_dashboard_is_idempotent_and_seeds_one_template_version(self) -> None:
        first = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        second = atlas_repository.atlas_dashboard(77, 42, "Пользователь")

        self.assertEqual(first["organization"]["id"], second["organization"]["id"])
        self.assertEqual(first["membership"]["role"], "owner")
        self.assertEqual(len(second["templates"]), 3)
        with connect() as con:
            self.assertEqual(
                con.execute("SELECT COUNT(*) FROM atlas_document_templates").fetchone()[0],
                3,
            )

    def test_overlay_character_binding_is_owned_scoped_and_uses_dynamic_catalog(self) -> None:
        character = storage.add_profile_character(77, 42, "Saul Goodman", "263345")
        storage.add_profile_character(77, 99, "Other Person", "777777")

        context = atlas_repository.atlas_set_overlay_character(
            77,
            42,
            character.id,
            server_code="phoenix-15",
            faction_code="fib",
            rank="Special Agent",
            screen_context_enabled=True,
        )

        self.assertTrue(context["selected_character"]["identity_verified"])
        self.assertEqual(context["selected_character"]["faction_code"], "fib")
        self.assertEqual(context["selected_character"]["rank"], "Special Agent")
        self.assertTrue(context["selected_character"]["screen_context_enabled"])
        preserved = atlas_repository.atlas_set_overlay_character(
            77,
            42,
            character.id,
            server_code="phoenix-15",
            faction_code="gov",
        )
        self.assertEqual(preserved["selected_character"]["rank"], "Special Agent")
        self.assertEqual(len(context["characters"]), 1)
        self.assertEqual(
            {item["code"] for item in context["catalog"]["factions"]},
            {"lspd", "lscsd", "fib", "gov", "sang", "ems", "wn"},
        )
        with self.assertRaisesRegex(ValueError, "atlas_overlay_character_not_owned"):
            atlas_repository.atlas_set_overlay_character(
                77,
                99,
                character.id,
                server_code="phoenix-15",
                faction_code="gov",
            )

    def test_overlay_binding_table_is_added_to_existing_database_without_data_loss(self) -> None:
        character = storage.add_profile_character(77, 42, "Saul Goodman", "263345")
        with connect() as con:
            con.execute("DROP TABLE atlas_character_bindings")
            con.commit()

        storage.init_db()

        with connect() as con:
            table = con.execute(
                "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
                ("atlas_character_bindings",),
            ).fetchone()
            preserved = con.execute(
                "SELECT nickname, static_id FROM profile_characters WHERE id = ?",
                (int(character.id),),
            ).fetchone()
        self.assertIsNotNone(table)
        self.assertEqual(tuple(preserved), ("Saul Goodman", "263345"))

    def test_answer_feedback_is_saved_and_bad_answer_is_excluded_from_memory(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        thread_id = atlas_repository.atlas_create_thread(organization_id, 42, "Проверка")
        atlas_repository.atlas_add_message(thread_id, "user", "Вопрос")
        message_id = atlas_repository.atlas_add_message(thread_id, "assistant", "Неточный ответ")

        feedback = atlas_repository.atlas_set_message_feedback(
            organization_id, 42, message_id, "bad"
        )
        thread = atlas_repository.atlas_thread_messages(organization_id, 42, thread_id)
        memory = atlas_repository.atlas_recent_chat_memory(organization_id, 42)

        self.assertEqual(feedback["rating"], "bad")
        self.assertEqual(thread["messages"][-1]["feedback_rating"], "bad")
        self.assertNotIn(message_id, {int(item["id"]) for item in memory})
        with self.assertRaisesRegex(ValueError, "atlas_feedback_message_not_found"):
            atlas_repository.atlas_set_message_feedback(
                organization_id, 99, message_id, "good"
            )

    def test_answer_provenance_round_trips_with_feedback_candidate(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        thread_id = atlas_repository.atlas_create_thread(
            organization_id,
            42,
            "Проверка выпуска",
            agent_id="atlas-claims",
        )
        atlas_repository.atlas_add_message(
            thread_id,
            "user",
            "Подготовь основу иска.",
            project_code="majestic-rp",
            server_code="phoenix-15",
            faction_code="gov",
        )
        message_id = atlas_repository.atlas_add_message(
            thread_id,
            "assistant",
            "Основа иска подготовлена по подтверждённой норме.",
            citations=[{"source_id": 41, "title": "Судебный кодекс"}],
            model="account/atlas-claims-v1",
            model_provider="together",
            model_release="atlas-claims-v1",
            project_code="majestic-rp",
            server_code="phoenix-15",
            faction_code="gov",
            latency_ms=87,
        )
        atlas_repository.atlas_set_message_feedback(organization_id, 42, message_id, "good")

        messages = atlas_repository.atlas_thread_messages(organization_id, 42, thread_id)["messages"]
        candidates = atlas_repository.atlas_training_candidates(
            organization_id=organization_id,
            project_code="majestic-rp",
            agent_id="atlas-claims",
        )

        assistant = messages[-1]
        self.assertEqual(assistant["model_provider"], "together")
        self.assertEqual(assistant["model_release"], "atlas-claims-v1")
        self.assertEqual(assistant["project_code"], "majestic-rp")
        self.assertEqual(assistant["server_code"], "phoenix-15")
        self.assertEqual(assistant["faction_code"], "gov")
        self.assertEqual(candidates[0]["model_provider"], "together")
        self.assertEqual(candidates[0]["answer_server_code"], "phoenix-15")

    def test_forum_attachment_ocr_is_review_gated_and_resets_on_source_revision(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Редактор")
        organization_id = int(dashboard["organization"]["id"])
        first = atlas_repository.atlas_upsert_synced_knowledge(
            organization_id,
            title="Судебный акт по делу №17",
            content="Проверяемая редакция материала форума с описанием приложенного судебного акта.",
            source_url="https://forum.majestic-rp.ru/threads/court-act.17/",
            server_code="phoenix-15",
            faction_code="gov",
            visibility_scope="server",
            feed_key="court-acts",
        )["source"]
        discovered = atlas_forum_attachment_repository.atlas_sync_forum_attachments(
            organization_id,
            int(first["id"]),
            (
                {
                    "url": "https://forum.majestic-rp.ru/attachments/court-act-17.100/",
                    "filename": "court-act-17.png",
                    "media_kind": "image",
                    "label": "Акт суда, лист 1",
                },
            ),
        )
        self.assertEqual(discovered[0]["status"], "discovered")
        attachment_id = int(discovered[0]["id"])
        atlas_forum_attachment_repository.atlas_forum_attachment_complete_ocr(
            attachment_id,
            content_sha256="a" * 64,
            mime_type="image/png",
            size_bytes=1234,
            storage_key="objects/aa/aa/" + "a" * 64,
            text="Проверяемая машинная расшифровка приложенного судебного акта.",
            engine="tesseract",
        )
        review = atlas_forum_attachment_repository.atlas_forum_attachment_review(
            organization_id,
            42,
            attachment_id,
            approve=True,
        )
        self.assertEqual(review["status"], "approved")
        self.assertGreater(int(review["knowledge_source_id"]), 0)
        derivative = atlas_repository.atlas_knowledge_source(
            int(review["knowledge_source_id"])
        )
        self.assertEqual(derivative["status"], "pending")
        self.assertIn("машинная расшифровка", derivative["content_text"])
        # The parent source remains the forum post; OCR has no path into the
        # parent text. The reviewed derivative is separately citable.
        self.assertNotIn("машинная расшифровка", first["content_text"])

        revised = atlas_repository.atlas_upsert_synced_knowledge(
            organization_id,
            title="Судебный акт по делу №17",
            content="Новая проверяемая редакция материала форума с изменённым описанием судебного акта.",
            source_url="https://forum.majestic-rp.ru/threads/court-act.17/",
            server_code="phoenix-15",
            faction_code="gov",
            visibility_scope="server",
            feed_key="court-acts",
        )["source"]
        reset = atlas_forum_attachment_repository.atlas_sync_forum_attachments(
            organization_id,
            int(revised["id"]),
            (
                {
                    "url": "https://forum.majestic-rp.ru/attachments/court-act-17.100/",
                    "filename": "court-act-17.png",
                    "media_kind": "image",
                },
            ),
        )[0]
        self.assertEqual(reset["status"], "discovered")
        self.assertIsNone(reset["ocr_text"])
        self.assertIsNone(reset["reviewed_at"])
        self.assertIsNone(reset["knowledge_source_id"])
        self.assertEqual(
            atlas_repository.atlas_knowledge_source(int(derivative["id"]))["status"],
            "archived",
        )

    def test_federation_migration_preserves_archived_source_scope(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Редактор")
        source = atlas_repository.atlas_add_knowledge(
            int(dashboard["organization"]["id"]),
            42,
            title="Исторический общий регламент",
            content="Этот регламент остаётся историческим материалом после миграции Atlas.",
            visibility_scope="global",
        )
        migration_key = "migration:atlas-federation:2026-08-31-v2"
        with connect() as con:
            con.execute(
                """
                UPDATE atlas_knowledge_sources
                SET project_code = '', federation_scope = 'workspace', status = 'archived'
                WHERE id = ?
                """,
                (int(source["id"]),),
            )
            con.execute("DELETE FROM meta WHERE key = ?", (migration_key,))
            con.commit()

        storage.init_db()

        with connect() as con:
            restored = con.execute(
                """
                SELECT project_code, federation_scope, status
                FROM atlas_knowledge_sources WHERE id = ?
                """,
                (int(source["id"]),),
            ).fetchone()
        self.assertEqual(tuple(restored), ("majestic-rp", "project", "archived"))

    def test_finetuning_candidates_require_good_feedback_and_pair_last_user_message(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        thread_id = atlas_repository.atlas_create_thread(organization_id, 42, "Обучение")
        atlas_repository.atlas_add_message(thread_id, "user", "Первый вопрос")
        rejected_id = atlas_repository.atlas_add_message(
            thread_id, "assistant", "Ответ без положительной оценки"
        )
        atlas_repository.atlas_set_message_feedback(
            organization_id, 42, rejected_id, "bad"
        )
        user_id = atlas_repository.atlas_add_message(thread_id, "user", "Точный вопрос")
        accepted_id = atlas_repository.atlas_add_message(
            thread_id,
            "assistant",
            "Точный и полезный ответ со ссылками.",
            citations=[{"title": "Уголовный кодекс"}],
            model="atlas-test",
        )
        atlas_repository.atlas_set_message_feedback(
            organization_id, 42, accepted_id, "good", comment="Проверено"
        )

        candidates = atlas_repository.atlas_training_candidates(
            organization_id=organization_id
        )

        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["assistant_message_id"], accepted_id)
        self.assertEqual(candidates[0]["user_message_id"], user_id)
        self.assertEqual(candidates[0]["user_text"], "Точный вопрос")
        self.assertEqual(candidates[0]["citations"][0]["title"], "Уголовный кодекс")

    def test_finetuning_candidates_are_filtered_by_project_and_agent(self) -> None:
        first = atlas_repository.atlas_dashboard(77, 42, "Majestic редактор")
        first_id = int(first["organization"]["id"])
        first_thread = atlas_repository.atlas_create_thread(
            first_id, 42, "Общий", agent_id="atlas-tvr-a"
        )
        atlas_repository.atlas_add_message(first_thread, "user", "Вопрос первого проекта")
        first_message = atlas_repository.atlas_add_message(
            first_thread, "assistant", "Проверенный и полезный ответ первого проекта."
        )
        atlas_repository.atlas_set_message_feedback(first_id, 42, first_message, "good")

        atlas_repository.atlas_upsert_project(42, code="project-b", name="Project B")
        server = atlas_repository.atlas_upsert_server(
            42,
            code="project-b-15",
            name="Phoenix",
            number=15,
            project_code="project-b",
        )
        second = atlas_repository.atlas_create_organization(
            77,
            42,
            name="Project B LSPD",
            owner_user_id=84,
            server_code=server["code"],
            faction_code="lspd",
        )
        second_id = int(second["id"])
        second_thread = atlas_repository.atlas_create_thread(
            second_id, 84, "Иск", agent_id="atlas-claims"
        )
        atlas_repository.atlas_add_message(second_thread, "user", "Вопрос второго проекта")
        second_message = atlas_repository.atlas_add_message(
            second_thread, "assistant", "Проверенный и полезный ответ второго проекта."
        )
        atlas_repository.atlas_set_message_feedback(second_id, 84, second_message, "good")

        majestic = atlas_repository.atlas_training_candidates(
            project_code="majestic-rp", agent_id="atlas-tvr-a"
        )
        second_project = atlas_repository.atlas_training_candidates(
            project_code="project-b", agent_id="atlas-claims"
        )
        mismatched = atlas_repository.atlas_training_candidates(
            organization_id=second_id,
            project_code="majestic-rp",
        )

        self.assertEqual([item["assistant_message_id"] for item in majestic], [first_message])
        self.assertEqual([item["project_code"] for item in majestic], ["majestic-rp"])
        self.assertEqual([item["assistant_message_id"] for item in second_project], [second_message])
        self.assertEqual([item["project_code"] for item in second_project], ["project-b"])
        self.assertEqual(mismatched, [])

    def test_agent_threads_have_isolated_memory_lanes(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        general = atlas_repository.atlas_create_thread(
            organization_id, 42, "Общий", agent_id="atlas-tvr-a"
        )
        claims = atlas_repository.atlas_create_thread(
            organization_id, 42, "Иск", agent_id="atlas-claims"
        )
        atlas_repository.atlas_add_message(general, "user", "Личная общая заметка")
        atlas_repository.atlas_add_message(claims, "user", "Факты искового дела")

        memory = atlas_repository.atlas_recent_chat_memory(
            organization_id, 42, agent_id="atlas-claims"
        )
        stored = atlas_repository.atlas_thread_messages(organization_id, 42, claims)

        self.assertEqual(stored["thread"]["agent_id"], "atlas-claims")
        self.assertEqual([item["content_text"] for item in memory], ["Факты искового дела"])
        self.assertNotIn("Личная общая заметка", str(memory))

    def test_onboarding_and_document_are_persisted(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        membership = atlas_repository.atlas_update_onboarding(
            organization_id,
            42,
            step=4,
            profile={"agency": "SGL", "position": "Адвокат"},
        )
        report_template = next(
            item for item in dashboard["templates"] if item["code"] == "incident-report"
        )
        document = atlas_repository.atlas_create_document(
            organization_id,
            42,
            title="Рапорт №1",
            template_id=int(report_template["id"]),
            fields={
                "date": "2026-08-22",
                "location": "Phoenix",
                "participants": "Пользователь",
                "facts": "Проверено",
                "actions": "Материал сохранён",
            },
        )

        self.assertEqual(membership["onboarding_step"], 4)
        self.assertEqual(membership["profile"]["agency"], "SGL")
        self.assertEqual(document["title"], "Рапорт №1")
        timeline = atlas_repository.atlas_timeline_events(organization_id)
        self.assertEqual(timeline[0]["event_kind"], "document")
        self.assertEqual(timeline[0]["source_id"], str(document["id"]))
        snapshot = atlas_repository.atlas_admin_snapshot(77)
        self.assertEqual(snapshot["totals"]["documents"], 1)
        self.assertEqual(snapshot["recent_events"][0]["event_type"], "document_created")
        with self.assertRaisesRegex(ValueError, "atlas_document_template_forbidden"):
            atlas_repository.atlas_create_document(
                organization_id,
                42,
                title="Документ с чужим шаблоном",
                template_id=999999,
                fields={},
            )

    def test_continuity_timeline_is_idempotent_and_space_isolated(self) -> None:
        own = atlas_repository.atlas_dashboard(77, 42, "Первый")
        foreign = atlas_repository.atlas_dashboard(88, 84, "Второй")
        organization_id = int(own["organization"]["id"])
        first = atlas_repository.atlas_create_timeline_event(
            organization_id,
            42,
            title="Конфликт у штаба",
            summary="Сохранён исходный контекст события.",
            event_kind="incident",
            importance="important",
            dedupe_key="incident:headquarters:1",
        )
        repeated = atlas_repository.atlas_create_timeline_event(
            organization_id,
            42,
            title="Повторная отправка",
            dedupe_key="incident:headquarters:1",
        )

        self.assertEqual(first["id"], repeated["id"])
        self.assertEqual(atlas_repository.atlas_timeline_summary(organization_id)["attention"], 1)
        self.assertEqual(len(atlas_repository.atlas_timeline_events(organization_id)), 1)
        self.assertEqual(
            atlas_repository.atlas_timeline_events(int(foreign["organization"]["id"])),
            [],
        )
        with self.assertRaisesRegex(ValueError, "atlas_timeline_forbidden"):
            atlas_repository.atlas_create_timeline_event(
                organization_id,
                999,
                title="Чужое событие",
            )

    def test_continuity_graph_links_existing_entities_without_duplicates(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        event = atlas_repository.atlas_create_timeline_event(
            organization_id,
            42,
            title="Материал получен",
            event_kind="activity",
        )
        document = atlas_repository.atlas_create_document(
            organization_id,
            42,
            title="Сводка события",
            template_id=None,
            fields={},
        )
        first = atlas_repository.atlas_link_entities(
            organization_id,
            42,
            source_type="timeline_event",
            source_id=event["id"],
            relation="produced",
            target_type="document",
            target_id=document["id"],
        )
        repeated = atlas_repository.atlas_link_entities(
            organization_id,
            42,
            source_type="timeline_event",
            source_id=event["id"],
            relation="produced",
            target_type="document",
            target_id=document["id"],
        )

        self.assertEqual(first["id"], repeated["id"])
        links = atlas_repository.atlas_entity_links(
            organization_id, "document", document["id"]
        )
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0]["source_id"], str(event["id"]))

        with self.assertRaisesRegex(ValueError, "atlas_entity_target_not_found"):
            atlas_repository.atlas_link_entities(
                organization_id,
                42,
                source_type="timeline_event",
                source_id=event["id"],
                relation="produced",
                target_type="document",
                target_id=999999,
            )
        with self.assertRaisesRegex(ValueError, "atlas_entity_type_invalid"):
            atlas_repository.atlas_link_entities(
                organization_id,
                42,
                source_type="invented",
                source_id=event["id"],
                relation="related_to",
                target_type="document",
                target_id=document["id"],
            )

    def test_timeline_update_uses_versions_and_page_cursor(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        created = [
            atlas_repository.atlas_create_timeline_event(
                organization_id,
                42,
                title=f"Событие {index}",
                occurred_at="2026-08-22T12:00:00+00:00",
            )
            for index in range(3)
        ]

        first_page = atlas_repository.atlas_timeline_page(organization_id, limit=2)
        second_page = atlas_repository.atlas_timeline_page(
            organization_id,
            limit=2,
            cursor=first_page["next_cursor"],
        )
        self.assertEqual([item["id"] for item in first_page["items"]], [created[2]["id"], created[1]["id"]])
        self.assertEqual([item["id"] for item in second_page["items"]], [created[0]["id"]])
        self.assertIsNone(second_page["next_cursor"])

        resolved = atlas_repository.atlas_update_timeline_event(
            organization_id,
            42,
            int(created[0]["id"]),
            status="resolved",
            summary="Результат сохранён.",
            expected_version=1,
        )
        self.assertEqual(resolved["version"], 2)
        self.assertEqual(resolved["status"], "resolved")
        self.assertIsNotNone(resolved["resolved_at"])
        with self.assertRaisesRegex(ValueError, "atlas_timeline_version_conflict"):
            atlas_repository.atlas_update_timeline_event(
                organization_id,
                42,
                int(created[0]["id"]),
                status="active",
                expected_version=1,
            )

        foreign = atlas_repository.atlas_dashboard(88, 84, "Чужой")
        with self.assertRaisesRegex(ValueError, "atlas_timeline_source_not_found"):
            atlas_repository.atlas_create_timeline_event(
                organization_id,
                42,
                title="Чужой источник",
                source_type="document",
                source_id=atlas_repository.atlas_create_document(
                    int(foreign["organization"]["id"]),
                    84,
                    title="Чужой документ",
                    template_id=None,
                    fields={},
                )["id"],
            )

    def test_admin_snapshot_does_not_mix_guilds(self) -> None:
        own = atlas_repository.atlas_dashboard(77, 42, "Первый")
        foreign = atlas_repository.atlas_dashboard(88, 84, "Второй")
        atlas_repository.atlas_create_document(
            int(own["organization"]["id"]), 42, title="Свой", template_id=None, fields={}
        )
        atlas_repository.atlas_create_document(
            int(foreign["organization"]["id"]), 84, title="Чужой", template_id=None, fields={}
        )

        snapshot = atlas_repository.atlas_admin_snapshot(77)
        self.assertEqual(snapshot["totals"]["organizations"], 1)
        self.assertEqual(snapshot["totals"]["documents"], 1)
        self.assertEqual({item["summary"] for item in snapshot["recent_events"]}, {"Создан документ «Свой»"})

    def test_knowledge_requires_editor_membership_and_deduplicates(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        first = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Регламент",
            content="Проверенный текст внутреннего регламента длиной больше двадцати символов.",
            source_kind="regulation",
        )
        second = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Регламент — новая подпись",
            content="Проверенный текст внутреннего регламента длиной больше двадцати символов.",
            source_kind="regulation",
        )

        self.assertEqual(first["id"], second["id"])
        with self.assertRaisesRegex(ValueError, "atlas_knowledge_forbidden"):
            atlas_repository.atlas_add_knowledge(
                organization_id,
                999,
                title="Чужой источник",
                content="Этот пользователь не состоит в выбранной организации Atlas.",
            )
        rebuild_sources = atlas_repository.atlas_indexable_knowledge_sources()
        self.assertEqual([item["id"] for item in rebuild_sources], [first["id"]])
        self.assertEqual(rebuild_sources[0]["content_text"], first["content_text"])
        timeline = atlas_repository.atlas_timeline_events(organization_id)
        self.assertEqual(len(timeline), 1)
        self.assertEqual(timeline[0]["source_type"], "knowledge_source")
        self.assertEqual(timeline[0]["source_id"], str(first["id"]))

    def test_searchable_corpus_keeps_old_reference_documents_ahead_of_recent_noise(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        law = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Уголовный кодекс штата San Andreas",
            content="Старая, но действующая нормативная база с полным текстом статей.",
            visibility_scope="server",
        )
        for index in range(4):
            atlas_repository.atlas_add_knowledge(
                organization_id,
                42,
                title=f"Новость форума {index}",
                content=f"Свежий информационный материал номер {index}, не являющийся кодексом.",
                visibility_scope="server",
            )

        sources = atlas_repository.atlas_searchable_knowledge_sources(
            organization_id,
            server_code="phoenix-15",
            faction_code="lspd",
            query_terms=("уголов",),
            limit=2,
        )

        self.assertEqual(sources[0]["id"], law["id"])
        self.assertEqual(sources[0]["title"], "Уголовный кодекс штата San Andreas")

    def test_searchable_corpus_finds_old_neutral_thread_by_body_text(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        old_thread = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Рассмотрено — дело 001",
            content="Редкая формулировка: фиолетовый протокол действует при проверке документов.",
            visibility_scope="server",
        )
        with connect() as con:
            con.execute(
                "UPDATE atlas_knowledge_sources SET updated_at = ? WHERE id = ?",
                ("2000-01-01T00:00:00+00:00", int(old_thread["id"])),
            )
        for index in range(365):
            atlas_repository.atlas_add_knowledge(
                organization_id,
                42,
                title=f"Новость форума {index}",
                content=f"Свежий общий материал без искомого термина, запись {index}.",
                visibility_scope="server",
            )

        sources = atlas_repository.atlas_searchable_knowledge_sources(
            organization_id,
            server_code="phoenix-15",
            faction_code="lspd",
            query_terms=("фиолетовый",),
            limit=5,
        )

        self.assertIn(int(old_thread["id"]), {int(item["id"]) for item in sources})

    def test_searchable_corpus_prioritizes_old_exact_identifier_hits(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        old_complaint = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Рассмотрено — дело 001",
            content="Жалоба на игрока со статиком 228392 и приложенными доказательствами.",
            visibility_scope="server",
        )
        with connect() as con:
            con.execute(
                "UPDATE atlas_knowledge_sources SET updated_at = ? WHERE id = ?",
                ("2000-01-01T00:00:00+00:00", int(old_complaint["id"])),
            )
        for index in range(610):
            atlas_repository.atlas_add_knowledge(
                organization_id,
                42,
                title=f"Рассмотрено — дело {index + 100}",
                content=f"Свежая жалоба без нужного идентификатора, запись {index}.",
                visibility_scope="server",
            )

        sources = atlas_repository.atlas_searchable_knowledge_sources(
            organization_id,
            server_code="phoenix-15",
            faction_code="lspd",
            query_terms=("228392",),
            limit=5,
        )

        self.assertIn(int(old_complaint["id"]), {int(item["id"]) for item in sources})

    def test_searchable_corpus_ranks_multi_term_cases_and_reserves_reference_sources(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        relevant_case = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Рассмотрено — дело 001",
            content=(
                "Проверка ареста и порядок обжалования подробно разобраны "
                "в материалах дела."
            ),
            visibility_scope="server",
        )
        law = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="Уголовный кодекс штата San Andreas",
            content="Нормативная база с действующими составами и процедурами.",
            visibility_scope="server",
        )
        for index in range(610):
            single_term = ("проверка", "арест", "обжалование")[index % 3]
            atlas_repository.atlas_add_knowledge(
                organization_id,
                42,
                title=f"Новая тема форума {index}",
                content=f"Краткое упоминание: {single_term}. Материал обновлён.",
                visibility_scope="server",
            )

        sources = atlas_repository.atlas_searchable_knowledge_sources(
            organization_id,
            server_code="phoenix-15",
            faction_code="lspd",
            query_terms=("проверка", "арест", "обжалование"),
            limit=40,
        )
        source_ids = {int(item["id"]) for item in sources}

        self.assertIn(int(relevant_case["id"]), source_ids)
        self.assertIn(int(law["id"]), source_ids)

    def test_knowledge_is_separated_by_server_and_faction(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        text = "Один и тот же общий текст может действовать в разных государственных структурах."
        lspd = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="LSPD",
            content=text,
            server_code="phoenix-15",
            faction_code="lspd",
        )
        gov = atlas_repository.atlas_add_knowledge(
            organization_id,
            42,
            title="GOV",
            content=text,
            server_code="phoenix-15",
            faction_code="gov",
        )

        self.assertNotEqual(lspd["id"], gov["id"])
        self.assertEqual(
            [item["id"] for item in atlas_repository.atlas_knowledge_sources(
                organization_id, server_code="phoenix-15", faction_code="gov"
            )],
            [gov["id"]],
        )
        with self.assertRaisesRegex(ValueError, "atlas_faction_invalid"):
            atlas_repository.atlas_normalize_scope("phoenix-15", "unknown")

    def test_knowledge_visibility_combines_shared_and_isolates_private_data(self) -> None:
        owner = atlas_repository.atlas_dashboard(77, 42, "Владелец")
        guest = atlas_repository.atlas_dashboard(77, 84, "Другая организация")
        owner_id = int(owner["organization"]["id"])
        guest_id = int(guest["organization"]["id"])

        created = {}
        for scope, faction in (
            ("global", "lspd"),
            ("server", "lspd"),
            ("faction", "lspd"),
            ("workspace", "lspd"),
        ):
            created[scope] = atlas_repository.atlas_add_knowledge(
                owner_id,
                42,
                title=scope,
                content=f"Проверенный материал области {scope}, достаточно длинный для Atlas.",
                server_code="phoenix-15",
                faction_code=faction,
                visibility_scope=scope,
            )

        gov_ids = {
            int(item["id"])
            for item in atlas_repository.atlas_knowledge_sources(
                guest_id,
                server_code="phoenix-15",
                faction_code="gov",
            )
        }
        self.assertEqual(
            gov_ids,
            {int(created["global"]["id"]), int(created["server"]["id"])},
        )

        guest_lspd_ids = {
            int(item["id"])
            for item in atlas_repository.atlas_knowledge_sources(
                guest_id,
                server_code="phoenix-15",
                faction_code="lspd",
            )
        }
        self.assertIn(int(created["faction"]["id"]), guest_lspd_ids)
        self.assertNotIn(int(created["workspace"]["id"]), guest_lspd_ids)

        owner_lspd_ids = {
            int(item["id"])
            for item in atlas_repository.atlas_knowledge_sources(
                owner_id,
                server_code="phoenix-15",
                faction_code="lspd",
            )
        }
        self.assertEqual(owner_lspd_ids, {int(item["id"]) for item in created.values()})
        with self.assertRaisesRegex(ValueError, "atlas_knowledge_scope_invalid"):
            atlas_repository.atlas_add_knowledge(
                owner_id,
                42,
                title="Ошибка",
                content="Материал с неизвестной областью доступа не должен сохраниться.",
                visibility_scope="unknown",
            )

    def test_projects_keep_legal_corpora_and_workspaces_isolated(self) -> None:
        first = atlas_repository.atlas_dashboard(77, 42, "Majestic редактор")
        first_id = int(first["organization"]["id"])
        atlas_repository.atlas_upsert_project(
            42,
            code="project-b",
            name="Project B",
        )
        second_server = atlas_repository.atlas_upsert_server(
            42,
            code="project-b-15",
            name="Phoenix",
            number=15,
            project_code="project-b",
        )
        second = atlas_repository.atlas_create_organization(
            77,
            42,
            name="Project B LSPD",
            owner_user_id=84,
            server_code=second_server["code"],
            faction_code="lspd",
        )
        second_id = int(second["id"])
        majestic = atlas_repository.atlas_add_knowledge(
            first_id,
            42,
            title="Закон Majestic",
            content="Проверенная норма Majestic RP, применимая только в первом проекте.",
            visibility_scope="global",
        )
        project_b = atlas_repository.atlas_add_knowledge(
            second_id,
            84,
            title="Закон Project B",
            content="Проверенная норма второго проекта, не применимая в Majestic RP.",
            server_code="project-b-15",
            faction_code="lspd",
            visibility_scope="global",
        )
        platform = atlas_repository.atlas_add_knowledge(
            second_id,
            84,
            title="Платформенная политика Atlas",
            content="Единая техническая политика Atlas, явно опубликованная для всех проектов.",
            server_code="project-b-15",
            faction_code="lspd",
            visibility_scope="global",
            federation_scope="platform",
        )

        first_visible = {
            int(item["id"])
            for item in atlas_repository.atlas_searchable_knowledge_sources(
                first_id,
                server_code="phoenix-15",
                faction_code="lspd",
            )
        }
        second_visible = {
            int(item["id"])
            for item in atlas_repository.atlas_searchable_knowledge_sources(
                second_id,
                server_code="project-b-15",
                faction_code="lspd",
            )
        }

        self.assertEqual(majestic["federation_scope"], "project")
        self.assertEqual(project_b["federation_scope"], "project")
        self.assertEqual(platform["federation_scope"], "platform")
        self.assertIn(int(majestic["id"]), first_visible)
        self.assertIn(int(platform["id"]), first_visible)
        self.assertNotIn(int(project_b["id"]), first_visible)
        self.assertIn(int(project_b["id"]), second_visible)
        self.assertIn(int(platform["id"]), second_visible)
        self.assertNotIn(int(majestic["id"]), second_visible)
        with self.assertRaisesRegex(ValueError, "atlas_organization_project_mismatch"):
            atlas_repository.atlas_searchable_knowledge_sources(
                first_id,
                server_code="project-b-15",
                faction_code="lspd",
            )

    def test_project_namespaces_personal_spaces_feeds_and_server_identity(self) -> None:
        atlas_repository.atlas_upsert_project(42, code="project-c", name="Project C")
        server = atlas_repository.atlas_upsert_server(
            42,
            code="project-c-15",
            name="Phoenix",
            number=15,
            project_code="project-c",
        )
        majestic_personal = atlas_repository.atlas_ensure_personal_space(77, 42, "Роберт")
        project_personal = atlas_repository.atlas_ensure_personal_space(
            77,
            42,
            "Роберт",
            project_code="project-c",
        )
        majestic_feed = atlas_repository.atlas_ensure_forum_feed(
            77,
            feed_key="laws",
            root_url="https://forum.majestic-rp.ru/forums/laws/",
        )
        project_feed = atlas_repository.atlas_ensure_forum_feed(
            77,
            feed_key="laws",
            root_url="https://forum.example.org/forums/laws/",
            server_code=server["code"],
            faction_code="lspd",
        )

        self.assertNotEqual(majestic_personal["organization"]["id"], project_personal["organization"]["id"])
        self.assertEqual(project_personal["organization"]["project_code"], "project-c")
        self.assertEqual(majestic_feed["feed_key"], "majestic-rp:laws")
        self.assertEqual(project_feed["feed_key"], "project-c:laws")
        self.assertEqual(
            atlas_repository.atlas_forum_sync_status(77, project_code="project-c")["id"],
            project_feed["id"],
        )
        with self.assertRaisesRegex(ValueError, "atlas_server_project_immutable"):
            atlas_repository.atlas_upsert_server(
                42,
                code="project-c-15",
                name="Moved",
                project_code="majestic-rp",
            )

    def test_admin_catalog_and_private_space_are_dynamic_and_isolated(self) -> None:
        server = atlas_repository.atlas_upsert_server(
            42,
            code="legacy-16",
            name="Legacy",
            number=16,
        )
        faction = atlas_repository.atlas_upsert_faction(
            42,
            code="ems",
            name="Emergency Medical Services",
            short_name="EMS",
        )
        private_space = atlas_repository.atlas_create_organization(
            77,
            42,
            name="Медицинский контур",
            owner_user_id=84,
            server_code=server["code"],
            faction_code=faction["code"],
        )
        selected = atlas_repository.atlas_dashboard(
            77,
            84,
            "Врач",
            organization_id=int(private_space["id"]),
        )
        outsider = atlas_repository.atlas_dashboard(77, 99, "Посторонний")

        self.assertEqual(selected["organization"]["id"], private_space["id"])
        self.assertEqual(selected["organization"]["branding"]["faction_code"], "ems")
        self.assertIn("legacy-16", {item["code"] for item in selected["catalog"]["servers"]})
        self.assertIn("ems", {item["code"] for item in selected["catalog"]["factions"]})
        self.assertNotIn(
            int(private_space["id"]),
            {int(item["id"]) for item in outsider["spaces"]},
        )

    def test_knowledge_taxonomy_is_saved_with_source(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Редактор")
        source = atlas_repository.atlas_add_knowledge(
            int(dashboard["organization"]["id"]),
            42,
            title="Правила сервера Phoenix",
            content=(
                "OOC правила сервера устанавливают требования администрации и наказания "
                "за нарушения игрового процесса."
            ),
            source_kind="url",
        )

        self.assertEqual(source["metadata"]["taxonomy"]["domain"], "ooc")
        self.assertEqual(source["metadata"]["taxonomy"]["corpus_kind"], "server_rule")

    def test_taxonomy_migration_preserves_canonical_source_and_rebuilds_only_index(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Редактор")
        source = atlas_repository.atlas_add_knowledge(
            int(dashboard["organization"]["id"]),
            42,
            title="Сохранённый кодекс",
            content="Полная сохранённая редакция кодекса длиной больше двадцати символов.",
        )
        atlas_repository.atlas_mark_knowledge_indexed(int(source["id"]), point_id="old-point")
        with connect() as con:
            con.execute(
                "UPDATE atlas_knowledge_sources SET metadata_json = '{}' WHERE id = ?",
                (int(source["id"]),),
            )
            con.execute("DELETE FROM meta WHERE key = ?", ("migration:atlas-taxonomy:2026-08-09-v1",))
            con.commit()

        storage.init_db()

        with connect() as con:
            stored = con.execute(
                "SELECT * FROM atlas_knowledge_sources WHERE id = ?",
                (int(source["id"]),),
            ).fetchone()
        self.assertIsNotNone(stored)
        self.assertEqual(stored["content_text"], source["content_text"])
        self.assertEqual(stored["status"], "pending")
        self.assertIsNone(stored["qdrant_point_id"])
        visible = atlas_repository.atlas_knowledge_sources(
            int(dashboard["organization"]["id"])
        )[0]
        self.assertEqual(visible["metadata"]["taxonomy"]["corpus_kind"], "law")

    def test_chat_history_is_ordered_and_private_between_users(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        first = atlas_repository.atlas_create_thread(organization_id, 42, "Первый чат")
        second = atlas_repository.atlas_create_thread(organization_id, 42, "Второй чат")
        foreign = atlas_repository.atlas_create_thread(organization_id, 999, "Чужой чат")
        atlas_repository.atlas_add_message(first, "user", "Подготовь речь о реформе")
        atlas_repository.atlas_add_message(first, "assistant", "Начнём с правовой основы.")
        atlas_repository.atlas_add_message(second, "user", "Предпочитаю спокойный официальный стиль")
        atlas_repository.atlas_add_message(foreign, "user", "Секрет другого пользователя")

        threads = atlas_repository.atlas_threads(organization_id, 42)
        first_history = atlas_repository.atlas_thread_messages(
            organization_id,
            42,
            first,
        )
        memory = atlas_repository.atlas_recent_chat_memory(
            organization_id,
            42,
            exclude_thread_id=first,
        )

        self.assertEqual({item["id"] for item in threads}, {first, second})
        self.assertEqual(
            [item["role"] for item in first_history["messages"]],
            ["user", "assistant"],
        )
        self.assertEqual([item["thread_title"] for item in memory], ["Второй чат"])
        self.assertNotIn("Секрет другого пользователя", str(memory))
        with self.assertRaisesRegex(ValueError, "atlas_thread_not_found"):
            atlas_repository.atlas_thread_messages(organization_id, 42, foreign)

    def test_discord_thread_binding_and_atlas_audit_are_durable(self) -> None:
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Пользователь")
        organization_id = int(dashboard["organization"]["id"])
        thread_id = atlas_repository.atlas_create_thread(organization_id, 42, "Задержание")
        bound = atlas_repository.atlas_bind_discord_thread(
            discord_thread_id=9001,
            guild_id=77,
            parent_channel_id=1494602485485277194,
            organization_id=organization_id,
            atlas_thread_id=thread_id,
            owner_user_id=42,
        )
        atlas_repository.atlas_record_event(
            organization_id,
            42,
            "ai_answer_created",
            "Atlas ответил в Discord",
            target_type="discord_thread",
            target_id=9001,
        )

        self.assertEqual(bound["atlas_thread_id"], thread_id)
        self.assertEqual(atlas_repository.atlas_discord_thread(9001)["owner_user_id"], 42)
        self.assertEqual(
            atlas_repository.atlas_admin_snapshot(77)["recent_events"][0]["target_id"],
            "9001",
        )


class AtlasAdminTokenWebTests(unittest.IsolatedAsyncioTestCase):
    async def test_only_administrator_can_grant_tokens_and_retry_is_idempotent(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-admin-token-web.db"
        storage.init_db()
        viewer = {
            "principal": ConsensusWebPrincipal(
                user_id=42,
                guild_id=77,
                display_name="Участник",
                csrf_token="token-csrf",
                member=SimpleNamespace(
                    id=42,
                    display_name="Участник",
                    guild_permissions=SimpleNamespace(administrator=False),
                    roles=[],
                ),
            )
        }

        async def authenticate(_request):
            return viewer["principal"], False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        payload = {
            "action": "grant",
            "user_id": "825331775857360906",
            "amount_tokens": 250_000,
            "reason": "Проверочное начисление администратора",
        }
        try:
            async with TestClient(TestServer(app)) as client:
                forbidden = await client.post(
                    "/api/admin/atlas/tokens",
                    json=payload,
                    headers={
                        "X-CSRF-Token": "token-csrf",
                        "X-Idempotency-Key": "grant-web-test-1",
                    },
                )
                viewer["principal"] = ConsensusWebPrincipal(
                    user_id=902235631952998410,
                    guild_id=77,
                    display_name="Администратор",
                    csrf_token="admin-csrf",
                    member=SimpleNamespace(
                        id=902235631952998410,
                        display_name="Администратор",
                        guild_permissions=SimpleNamespace(administrator=True),
                        roles=[],
                    ),
                )
                headers = {
                    "X-CSRF-Token": "admin-csrf",
                    "X-Idempotency-Key": "grant-web-test-1",
                }
                granted = await client.post(
                    "/api/admin/atlas/tokens",
                    json=payload,
                    headers=headers,
                )
                repeated = await client.post(
                    "/api/admin/atlas/tokens",
                    json=payload,
                    headers=headers,
                )
                lookup = await client.get(
                    "/api/admin/atlas/tokens?user_id=825331775857360906"
                )
                granted_body = await granted.json()
                repeated_body = await repeated.json()
                lookup_body = await lookup.json()

            self.assertEqual(forbidden.status, 403)
            self.assertEqual(granted.status, 201)
            self.assertEqual(repeated.status, 201)
            self.assertEqual(granted_body["entry"]["id"], repeated_body["entry"]["id"])
            self.assertEqual(lookup_body["summary"]["payg_balance_tokens"], 250_000)
            with connect() as con:
                self.assertEqual(
                    con.execute(
                        "SELECT COUNT(*) FROM bot_actions WHERE module = 'atlas_billing'"
                    ).fetchone()[0],
                        1,
                )
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()


class AtlasAITests(unittest.IsolatedAsyncioTestCase):
    @staticmethod
    def _project_rules_source() -> dict[str, object]:
        return {
            "id": 9_071,
            "organization_id": 77,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "global",
            "federation_scope": "project",
            "title": "Основные правила проекта",
            "source_kind": "forum",
            "source_url": "https://forum.majestic-rp.ru/threads/osnovnyye-pravila-proyekta.8036/",
            "content_text": (
                "Общее положение\n"
                "1.1 На проекте действует прецедентная система правил.\n"
                "Положение об аккаунте\n"
                "2.1 Максимальное количество разрешенных аккаунтов на одного человека — один.\n"
                "2.2 Запрещено передавать аккаунт 3-м лицам. | PermBan.\n"
                "2.2.1 Вложенное пояснение к передаче аккаунта.\n"
                "2.20 Условный соседний пункт, который не относится к передаче.\n"
                "2.3 Администрация не несет ответственности за аккаунт при взломе.\n"
                "Игровые чаты\n"
                "4.1 Текстовый и голосовой чат является исключительно IC чатом, где запрещено "
                "OOC-общение; OOC информация передается через /b, /fb, /gb и /cb. | Mute 30–90 "
                "минут / Demorgan 5 минут.\n"
                "4.3 Запрещено прямое оскорбление родственников. | HardBan 30–60 дней.\n"
                "Role Play процесс\n"
                "5.1 DM — прямое убийство, нанесение урона или стрельба без IC причины и IC диалога. "
                "| GunBan 8 часов / Demorgan 120 минут / WARN / Ban 3–30 дней.\n"
                "Исключение: IC диалог не обязателен при угрозе жизни, грубых оскорблениях, угоне "
                "транспортного средства и других прямо перечисленных ситуациях.\n"
                "5.2 DB — умышленный наезд транспортом.\n"
            ),
            "metadata": {
                "taxonomy": {
                    "domain": "ooc",
                    "corpus_kind": "server_rule",
                    "authority_scope": "project",
                }
            },
        }

    def test_answer_text_accepts_provider_content_variants(self) -> None:
        self.assertEqual(
            _answer_text(
                {
                    "choices": [{
                        "message": {
                            "content": [
                                {"type": "output_text", "text": {"value": "Готовый "}},
                                {"type": "text", "text": "ответ"},
                            ]
                        }
                    }]
                }
            ),
            "Готовый ответ",
        )
        self.assertEqual(
            _answer_text(
                {"choices": [{"message": {"content": None}, "text": "Резервный ответ"}]}
            ),
            "Резервный ответ",
        )

    def test_agent_registry_preserves_general_and_adds_specialists(self) -> None:
        catalog = atlas_agent_catalog()
        self.assertEqual(catalog[0]["id"], "atlas-tvr-a")
        self.assertIn("atlas-claims", {item["id"] for item in catalog})
        self.assertIn("исков", atlas_resolve_agent("atlas-claims").specialty.casefold())
        self.assertEqual(atlas_resolve_agent("atlas-claims").knowledge_domains, ("ic", "mixed"))
        self.assertEqual(atlas_resolve_agent("atlas-complaints").knowledge_domains, ("ooc", "mixed"))
        self.assertEqual(atlas_resolve_agent("atlas-complaints").training_lane, "ooc-complaints")
        self.assertIn(
            "Не переноси требования к доказательствам",
            atlas_resolve_agent("atlas-complaints").instruction,
        )
        self.assertIn(
            "Любые не названные пользователем",
            atlas_resolve_agent("atlas-complaints").instruction,
        )
        self.assertIn(
            "Не требуй конкретное наказание",
            atlas_resolve_agent("atlas-complaints").instruction,
        )
        with self.assertRaisesRegex(ValueError, "atlas_agent_invalid"):
            atlas_resolve_agent("unknown")

    def test_inflected_short_request_gets_the_strict_short_contract(self) -> None:
        question = "Определи нарушение и составь краткую жалобу."
        profile = _atlas_task_profile(question, mode="balanced")

        self.assertEqual(profile.intent, "drafting")
        self.assertEqual(profile.reasoning_effort, "medium")
        self.assertIn("жёсткий предел — 70 слов", _response_delivery_contract(profile, question))

    def test_default_answer_contract_is_compact_without_requesting_a_deep_analysis(self) -> None:
        question = "Что делать, если остановили?"
        profile = _atlas_task_profile(question, mode="balanced")

        self.assertIn("жёсткий предел — 60 слов", _response_delivery_contract(profile, question))

    def test_expensive_legacy_default_is_downgraded_to_economy_model(self) -> None:
        with patch.dict(
            os.environ,
            {"ATLAS_OPENROUTER_MODEL": "openai/gpt-5.4"},
        ):
            self.assertEqual(atlas_ai_config().chat_model, "openai/gpt-5-mini")

    def test_weak_legacy_default_is_upgraded_to_economy_model(self) -> None:
        with patch.dict(
            os.environ,
            {"ATLAS_OPENROUTER_MODEL": "openai/gpt-4.1-mini"},
        ):
            self.assertEqual(atlas_ai_config().chat_model, "openai/gpt-5-mini")

    def test_custom_atlas_model_is_preserved(self) -> None:
        with patch.dict(
            os.environ,
            {"ATLAS_OPENROUTER_MODEL": "custom/provider-model"},
        ):
            self.assertEqual(atlas_ai_config().chat_model, "custom/provider-model")

    def test_deprecated_direct_model_is_upgraded_to_grok_43(self) -> None:
        with patch.dict(
            os.environ,
            {"ATLAS_DIRECT_MODEL": "x-ai/grok-4.1-fast"},
        ):
            self.assertEqual(atlas_ai_config().direct_model, "x-ai/grok-4.3")

    def test_fine_tuned_route_is_scoped_and_keeps_base_fallback(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="base-key",
            openrouter_url="https://openrouter.test/chat/completions",
            chat_model="openai/base",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
            together_key="together-key",
            fine_tuned_model="cdnserver/atlas-general-v1",
            fine_tuned_enabled=True,
            fine_tuned_agents="atlas-tvr-a",
            fine_tuned_projects="majestic-rp",
        )
        primary, fallback, reason = atlas_model_route(
            config,
            agent=atlas_resolve_agent("atlas-tvr-a"),
            project_code="majestic-rp",
        )
        self.assertEqual((primary.provider, primary.model, primary.release), (
            "together", "cdnserver/atlas-general-v1", "fine-tuned",
        ))
        self.assertEqual(fallback.model if fallback else None, "openai/base")
        self.assertEqual(reason, "fine_tuned_rollout")

        special, no_fallback, special_reason = atlas_model_route(
            config,
            agent=atlas_resolve_agent("atlas-tvr-a"),
            project_code="majestic-rp",
            direct_mode=True,
        )
        self.assertEqual(special.provider, "openrouter")
        self.assertEqual(no_fallback, None)
        self.assertEqual(special_reason, "special_route")

        isolated, isolated_fallback, isolated_reason = atlas_model_route(
            config,
            agent=atlas_resolve_agent("atlas-tvr-a"),
            project_code="another-project",
        )
        self.assertEqual(isolated.model, "openai/base")
        self.assertEqual(isolated_fallback, None)
        self.assertEqual(isolated_reason, "project_not_enrolled")

    def test_fine_tuned_rollout_registry_prefers_exact_project_and_agent(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="base-key",
            openrouter_url="https://openrouter.test/chat/completions",
            chat_model="openai/base",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
            together_key="together-key",
            fine_tuned_enabled=True,
            fine_tuned_rollouts=json.dumps([
                {"project_code": "*", "agent_id": "*", "model": "account/shared", "enabled": True},
                {
                    "project_code": "majestic-rp", "agent_id": "atlas-claims",
                    "model": "account/claims-v1", "release": "claims-v1", "enabled": True,
                },
            ]),
        )
        route, fallback, reason = atlas_model_route(
            config,
            agent=atlas_resolve_agent("atlas-claims"),
            project_code="majestic-rp",
        )
        self.assertEqual(route.model, "account/claims-v1")
        self.assertEqual(route.release, "claims-v1")
        self.assertEqual(fallback.model if fallback else None, "openai/base")
        self.assertEqual(reason, "fine_tuned_rollout")

    def test_atlas_2_is_text_only_and_strips_its_call_prefix(self) -> None:
        self.assertEqual(
            atlas_parse_text_mode("  Атлас 2, скажи прямо  "),
            ("скажи прямо", True),
        )
        self.assertEqual(
            atlas_parse_text_mode("Атлас 2, скажи прямо", latency_mode="overlay"),
            ("Атлас 2, скажи прямо", False),
        )

    async def test_atlas_2_routes_to_direct_model_and_reports_mode(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat/completions",
            chat_model="test/standard",
            direct_model="x-ai/test-direct",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        completion = AsyncMock(
            return_value={"choices": [{"message": {"content": "Прямой ответ"}}]}
        )
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search",
            AsyncMock(return_value=[]),
        ), patch("modules.atlas_ai._json_request", completion):
            result = await atlas_answer(77, "Атлас 2, скажи прямо")

        payload = completion.await_args.kwargs["payload"]
        self.assertEqual(payload["model"], "x-ai/test-direct")
        self.assertEqual(payload["messages"][-1]["content"], "скажи прямо")
        self.assertIn("без стилистической цензуры", payload["messages"][0]["content"])
        self.assertEqual(result["text_mode"], "atlas-2")

    async def test_retryable_fine_tuned_failure_falls_back_to_base_model(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="base-key",
            openrouter_url="https://openrouter.test/chat/completions",
            chat_model="openai/base",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
            together_key="together-key",
            together_url="https://together.test/v1/chat/completions",
            fine_tuned_model="cdnserver/atlas-general-v1",
            fine_tuned_enabled=True,
            fine_tuned_projects="majestic-rp",
        )
        calls: list[tuple[str, str]] = []

        async def complete(_method, url, *, payload, **_kwargs):
            calls.append((url, str(payload.get("model"))))
            if "together.test" in url:
                raise AtlasAIError("upstream_unavailable", "temporary", retryable=True)
            return {"choices": [{"message": {"content": "Базовый ответ"}}]}

        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search",
            AsyncMock(return_value=[]),
        ), patch("modules.atlas_ai._json_request", side_effect=complete):
            result = await atlas_answer(77, "Помоги составить короткую речь")

        self.assertEqual(
            calls,
            [
                ("https://together.test/v1/chat/completions", "cdnserver/atlas-general-v1"),
                ("https://openrouter.test/chat/completions", "openai/base"),
            ],
        )
        self.assertEqual(result["model"], "openai/base")
        self.assertEqual(result["model_provider"], "openrouter")

    def test_taxonomy_distinguishes_ic_ooc_charters_and_case_law(self) -> None:
        ooc = atlas_classify_knowledge(
            title="Правила сервера",
            content="OOC требования проекта и наказания администрации за нарушения.",
        )
        charter = atlas_classify_knowledge(
            title="Устав LSPD",
            content="Настоящий устав определяет полномочия сотрудников полиции и ранги.",
        )
        practice = atlas_classify_knowledge(
            title="Судебная практика",
            content="Решение суда по исковому заявлению и толкованию положений закона.",
        )

        self.assertEqual((ooc["domain"], ooc["corpus_kind"]), ("ooc", "server_rule"))
        self.assertEqual((charter["domain"], charter["corpus_kind"]), ("ic", "charter"))
        self.assertEqual(practice["corpus_kind"], "case_law")
        self.assertEqual(practice["authority_scope"], "court")

    def test_taxonomy_keeps_project_rules_ooc_when_they_describe_ic_mechanics(self) -> None:
        # Real project rules contain terms such as IC, DM and RP.  They are
        # still OOC regulations, rather than an IC/OOC mixed corpus.
        rules = atlas_classify_knowledge(
            title="Основные правила проекта",
            content=(
                "На проекте действует система наказаний. Текстовый и голосовой "
                "чат является IC, запрещены DM, DB, NonRP и нарушения RP процесса."
            ),
            source_url="https://forum.majestic-rp.ru/threads/osnovnyye-pravila-proyekta.8036/",
            source_kind="forum",
        )

        self.assertEqual((rules["domain"], rules["corpus_kind"]), ("ooc", "server_rule"))
        self.assertEqual(rules["authority_scope"], "project")

    def test_taxonomy_uses_document_title_before_incidental_body_terms(self) -> None:
        criminal_code = atlas_classify_knowledge(
            title="Уголовный Кодекс штата San Andreas",
            content=(
                "Суд рассматривает материалы дела. Прокурор может издать распоряжение, "
                "а участник вправе подать исковое заявление."
            ),
            source_kind="forum",
        )
        constitution = atlas_classify_knowledge(
            title="Конституция Штата San Andreas",
            content="Судебная практика и исковые заявления применяются с учетом Конституции.",
            source_kind="forum",
        )

        self.assertEqual((criminal_code["domain"], criminal_code["corpus_kind"]), ("ic", "law"))
        self.assertEqual((constitution["domain"], constitution["corpus_kind"]), ("ic", "law"))

    def test_aristotle_plan_adapts_to_legal_case(self) -> None:
        plan = atlas_research_plan("Составь позицию по иску и судебной практике")
        self.assertEqual(plan[0]["agent"], "Навигатор")
        self.assertEqual(plan[-1]["agent"], "Аристотель")
        self.assertIn("practice", {item["id"] for item in plan})

    def test_chunker_is_bounded_and_preserves_overlap(self) -> None:
        chunks = _chunks("слово " * 4000, size=1000, overlap=100)
        self.assertGreater(len(chunks), 2)
        self.assertLessEqual(len(chunks), 48)
        self.assertTrue(all(len(chunk) <= 1000 for chunk in chunks))

    def test_query_variants_expand_legal_abbreviation(self) -> None:
        variants = _atlas_query_variants("Что такое УК?")

        self.assertTrue(any("уголовный кодекс" in item.casefold() for item in variants))

    def test_query_variants_route_colloquial_crime_to_criminal_code(self) -> None:
        variants = _atlas_query_variants("Назови статью за убийство")

        self.assertTrue(
            any("уголовный кодекс штата san andreas" in item.casefold() for item in variants)
        )

    def test_query_variants_route_software_check_to_its_ooc_rules(self) -> None:
        variants = _atlas_query_variants("Можно ли использовать стороннее ПО и как проходит проверка?")

        self.assertTrue(
            any("правила проверки на стороннее по" in item.casefold() for item in variants)
        )

    def test_retrieval_refusal_detector_ignores_a_substantive_no_prohibition_answer(self) -> None:
        self.assertTrue(
            _atlas_answer_is_retrieval_refusal(
                "В текущей библиотеке точная статья не найдена, поэтому назвать её не могу."
            )
        )
        self.assertFalse(
            _atlas_answer_is_retrieval_refusal(
                "В статье 6.2 нет отдельного запрета на оказание первой помощи."
            )
        )
        self.assertTrue(
            _atlas_answer_is_retrieval_refusal(
                "В предоставленной мне библиотеке источников нет полного текста главы 16."
            )
        )
        self.assertTrue(
            _atlas_answer_is_retrieval_refusal(
                "Не могу точно сказать, потому что соответствующий фрагмент отсутствует."
            )
        )
        for variant in (
            "В этом контексте нет информации о статье.",
            "По запросу ничего не найдено.",
            "Atlas не знает ответа по этому вопросу.",
            "В релевантных источниках отсутствуют данные.",
            "Библиотека пока не содержит доступных названий.",
            "No relevant information in the knowledge base.",
            "Knowledge base does not contain relevant information.",
            "I couldn't find the answer in the provided context.",
            "I don't have enough information in the available sources.",
            "Контекст не содержит применимой нормы.",
        ):
            self.assertTrue(_atlas_answer_is_retrieval_refusal(variant), variant)
        self.assertFalse(
            _atlas_answer_is_retrieval_refusal(
                "В статье 6.2 нет отдельного запрета на оказание первой помощи."
            )
        )

    def test_retrieval_refusal_fallback_returns_exact_structured_evidence(self) -> None:
        prepared = SimpleNamespace(
            payload={"messages": [{"role": "user", "content": "Покажи статью 10.1 УК"}]},
            sources=[
                {
                    "structured": True,
                    "text": "10.1 Кража — тайное хищение чужого имущества.",
                    "reference": "article:10.1",
                    "pinpoints": ["статья 10.1"],
                }
            ]
        )

        answer = _grounded_refusal_fallback(prepared)

        self.assertIn("10.1 Кража", answer)
        self.assertIn("[1, статья 10.1]", answer)
        self.assertNotIn("информации нет", answer.casefold())

    def test_retrieval_refusal_does_not_return_unrequested_structured_clause(self) -> None:
        source = {
            "structured": True,
            "text": "1.1 На охраняемых территориях имеют право находиться все граждане.",
            "reference": "article:1.1",
            "pinpoints": ["статья 1.1"],
        }
        prepared = SimpleNamespace(
            intent="legal_analysis",
            payload={"messages": [{"role": "user", "content": "Нарушают ли законы мой внешний вид?"}]},
            sources=[source],
        )

        self.assertEqual(_grounded_refusal_fallback(prepared), "")
        self.assertEqual(_deterministic_exact_lookup(prepared), "")

    def test_routed_query_drops_off_route_retrieval_hits(self) -> None:
        road_code = {"title": "Дорожный Кодекс штата San Andreas", "source_id": 10}
        unrelated = {"title": "Закон о статусе охраняемых территорий", "source_id": 11}

        self.assertEqual(
            _atlas_relevant_sources("Можно ли эвакуировать автомобиль?", [unrelated, road_code]),
            [road_code],
        )
        self.assertEqual(
            _atlas_relevant_sources("Можно ли эвакуировать автомобиль?", [unrelated]),
            [],
        )

    def test_unrouted_query_discards_structured_neighbour_without_overlap(self) -> None:
        unrelated = {
            "title": "Закон о статусе охраняемых территорий",
            "source_id": 11,
            "structured": True,
            "reference": "article:1.1",
            "text": "1.1 На охраняемых территориях имеют право находиться все граждане.",
        }

        self.assertEqual(
            _atlas_relevant_sources("Нарушают ли законы мой внешний вид?", [unrelated]),
            [],
        )

    def test_unrouted_query_keeps_sources_with_lexical_evidence(self) -> None:
        sources = [
            {
                "title": "Уголовный кодекс",
                "source_id": 10,
                "text": "Кража имущества является преступлением и наказывается по закону.",
            },
            {
                "title": "Кодекс этики",
                "source_id": 11,
                "text": "Правила этики регулируют поведение государственных служащих.",
            },
        ]

        self.assertEqual(
            _atlas_relevant_sources("Как наказывается кража имущества?", sources),
            [sources[0]],
        )

    def test_explicit_structured_reference_survives_relevance_filter(self) -> None:
        requested = {
            "title": "Уголовный кодекс",
            "source_id": 10,
            "structured": True,
            "reference": "article:10.1",
            "text": "10.1 Кража — тайное хищение чужого имущества.",
        }

        self.assertEqual(
            _atlas_relevant_sources("Покажи статью 10.1 УК", [requested]),
            [requested],
        )

    def test_ic_legal_query_does_not_receive_an_ooc_rescue_variant(self) -> None:
        variants = _atlas_query_variants(
            "Меня задержали сотрудники LSPD, какие у меня права?"
        )

        self.assertTrue(any("IC законодательство" in item for item in variants))
        self.assertFalse(any("OOC правила" in item for item in variants))

    def test_corpus_abbreviations_follow_actual_atlas_documents(self) -> None:
        aliases = _atlas_corpus_abbreviations(
            [
                {"title": "Процессуальный Кодекс штата San Andreas"},
                {"title": "Административный кодекс штата San Andreas"},
                {"title": "Кодекс этики и служебного поведения"},
            ]
        )

        self.assertEqual(aliases["пк"], "Процессуальный Кодекс штата San Andreas")
        self.assertEqual(aliases["ак"], "Административный кодекс штата San Andreas")
        self.assertEqual(aliases["кэ"], "Кодекс этики и служебного поведения")
        self.assertNotIn("упк", aliases)

    def test_corpus_abbreviations_include_latin_organization_names(self) -> None:
        aliases = _atlas_corpus_abbreviations(
            [
                {"title": 'Закон "О статусе United States Secret Service"'},
                {"title": 'Закон "О статусе Federal Investigation Bureau"'},
                {"title": 'Закон "О Статусе San Andreas National Guard"'},
            ]
        )

        self.assertEqual(aliases["usss"], "United States Secret Service")
        self.assertEqual(aliases["fib"], "Federal Investigation Bureau")
        self.assertEqual(aliases["sang"], "San Andreas National Guard")
        self.assertNotIn("коап", aliases)

    def test_bad_answer_is_not_reused_inside_current_dialog(self) -> None:
        messages = _bounded_dialog_messages(
            [
                {"role": "user", "content_text": "Вопрос"},
                {"role": "assistant", "content_text": "Ошибка", "feedback_rating": "bad"},
                {"role": "user", "content_text": "Попробуй ещё раз"},
            ]
        )

        self.assertEqual([item["content"] for item in messages], ["Вопрос", "Попробуй ещё раз"])

    def test_followup_search_uses_only_one_relevant_dialog_anchor(self) -> None:
        dialog = _bounded_dialog_messages(
            [
                {"role": "user", "content_text": "Сначала обсуждали форму рапорта"},
                {"role": "assistant", "content_text": "Вот форма"},
                {"role": "user", "content_text": "Теперь обсуждаем задержание по статье 16.1"},
                {"role": "assistant", "content_text": "Проверяю норму"},
            ]
        )

        profile = _atlas_task_profile(
            "А теперь объясни её простыми словами",
            mode="balanced",
            dialog_messages=dialog,
        )

        self.assertTrue(profile.is_followup)
        self.assertIn("задержание по статье 16.1", profile.retrieval_query)
        self.assertNotIn("форму рапорта", profile.retrieval_query)

    def test_fresh_question_does_not_pollute_search_with_old_dialog(self) -> None:
        profile = _atlas_task_profile(
            "Что такое Уголовный кодекс?",
            mode="balanced",
            dialog_messages=[
                {"role": "user", "content": "Составь речь про выборы"},
                {"role": "assistant", "content": "Готовая речь"},
            ],
        )

        self.assertFalse(profile.is_followup)
        self.assertEqual(profile.retrieval_query, "Что такое Уголовный кодекс?")
        self.assertNotIn("выборы", profile.retrieval_query)

        second = _atlas_task_profile(
            "А меня задержали, что делать?",
            mode="balanced",
            dialog_messages=[{"role": "user", "content": "Составь речь про выборы"}],
        )
        self.assertFalse(second.is_followup)
        self.assertNotIn("выборы", second.retrieval_query)

    def test_exact_chapter_request_is_not_treated_as_freeform_drafting(self) -> None:
        profile = _atlas_task_profile(
            "Напиши мне 16 главу УК полностью",
            mode="balanced",
        )

        self.assertEqual(profile.intent, "exact_lookup")
        self.assertEqual(profile.depth, "deep")

    def test_corpus_abbreviation_is_understood_by_task_router(self) -> None:
        profile = _atlas_task_profile("Что такое УК?", mode="balanced")

        self.assertEqual(profile.intent, "legal_analysis")
        self.assertEqual(profile.depth, "quick")

    def test_concise_modifier_does_not_replace_legal_or_procedural_intent(self) -> None:
        detention = _atlas_task_profile(
            "Меня задержали сотрудники LSPD. Кратко: какие у меня права и что делать?",
            mode="balanced",
        )
        prosecutor = _atlas_task_profile(
            "Кратко объясни полномочия Генерального прокурора.",
            mode="balanced",
        )

        self.assertEqual(detention.intent, "procedural_advice")
        self.assertEqual(detention.depth, "quick")
        self.assertEqual(prosecutor.intent, "legal_analysis")
        self.assertEqual(prosecutor.depth, "quick")

    def test_plain_greeting_stays_social_instead_of_describing_the_interface(self) -> None:
        profile = _atlas_task_profile("Привет!", mode="balanced")

        self.assertEqual(profile.intent, "social")
        self.assertEqual(profile.depth, "quick")
        self.assertIn("обычное человеческое обращение", profile.response_brief)

    def test_contextual_drafting_reuses_the_described_situation(self) -> None:
        profile = _atlas_task_profile(
            "На основе уже описанной ситуации составь жалобу",
            mode="balanced",
            dialog_messages=[
                {"role": "user", "content": "Меня задержали без представления сотрудника"},
                {"role": "assistant", "content": "Проверяю процедуру"},
            ],
        )

        self.assertTrue(profile.is_followup)
        self.assertEqual(profile.intent, "drafting")
        self.assertIn("задержали без представления", profile.retrieval_query)

    def test_research_context_keeps_two_latest_user_corrections(self) -> None:
        context = _recent_user_dialog_context(
            [
                {"role": "user", "content": "Старая не относящаяся тема"},
                {"role": "assistant", "content": "Старый ответ"},
                {"role": "user", "content": "Сотрудник не представился"},
                {"role": "assistant", "content": "Первичный анализ"},
                {"role": "user", "content": "Но ордер выдала прокуратура"},
            ]
        )

        self.assertNotIn("Старая", context)
        self.assertIn("не представился", context)
        self.assertIn("ордер выдала прокуратура", context)

    def test_pinpoint_labels_are_conservative_and_support_exact_references(self) -> None:
        self.assertEqual(
            _atlas_pinpoint_labels({"reference": "chapter:16", "text": "16.1 Текст"}),
            ["глава 16"],
        )
        self.assertEqual(
            _atlas_pinpoint_labels(
                {
                    "text": (
                        "2.6 Порядок начала задержания.\n"
                        "Описание нормы.\n"
                        "2.11 Право задержанного на защиту."
                    )
                }
            ),
            ["статья 2.6", "статья 2.11"],
        )

    def test_cross_chat_memory_does_not_reuse_unrated_model_claims(self) -> None:
        context = _cross_chat_context(
            [
                {
                    "role": "user",
                    "thread_title": "Предпочтения",
                    "content_text": "Пиши официально и кратко",
                },
                {
                    "role": "assistant",
                    "thread_title": "Старый ответ",
                    "content_text": "Статья якобы разрешает обыск",
                },
                {
                    "role": "assistant",
                    "feedback_rating": "good",
                    "thread_title": "Подтверждённый ответ",
                    "content_text": "Пользователь подтвердил этот удачный шаблон",
                },
            ]
        )

        self.assertIn("Пиши официально", context)
        self.assertNotIn("якобы разрешает", context)
        self.assertIn("удачный шаблон", context)

    def test_dialog_keeps_initial_user_brief_and_recent_turns(self) -> None:
        history = [{"role": "user", "content_text": "Главная цель: защитить клиента"}]
        for index in range(40):
            history.append(
                {
                    "role": "assistant" if index % 2 else "user",
                    "content_text": f"Сообщение {index}",
                }
            )

        messages = _bounded_dialog_messages(history, max_messages=8, max_chars=2000)

        self.assertEqual(len(messages), 8)
        self.assertEqual(messages[0]["content"], "Главная цель: защитить клиента")
        self.assertEqual(messages[-1]["content"], "Сообщение 39")

    async def test_hybrid_search_finds_saved_source_when_qdrant_returns_nothing(self) -> None:
        source = {
            "id": 91,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный кодекс",
            "content_text": (
                "Уголовный кодекс устанавливает основания ответственности, "
                "виды преступлений и применяемые наказания."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.1/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(77, "Что такое УК?", expanded=True)

        self.assertTrue(result)
        self.assertEqual(result[0]["source_id"], 91)
        self.assertEqual(result[0]["title"], "Уголовный кодекс")

    async def test_search_returns_explicitly_requested_legal_chapter(self) -> None:
        source = {
            "id": 93,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "Глава 15.\nПредыдущие нормы.\n"
                "Глава 16.\nПреступления против правосудия.\nСтатья 16.1. Точный текст.\n"
                "Глава 17.\nСледующие нормы."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.3/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(77, "Напиши мне 16 главу УК", expanded=True)

        self.assertTrue(result[0]["structured"])
        self.assertIn("Глава 16", result[0]["text"])
        self.assertIn("Статья 16.1", result[0]["text"])
        self.assertNotIn("Глава 17", result[0]["text"])

    async def test_search_routes_short_murder_question_to_criminal_code(self) -> None:
        criminal_code = {
            "id": 9_310,
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "Глава 6. Преступления против жизни.\n"
                "6.2 (F/R) Убийство, то есть умышленное причинение смерти другому человеку. "
                "Приоритет розыска — 4. Наказание: до 40 месяцев лишения свободы.\n"
                "6.3 Тяжкое убийство двух или более лиц."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.9310/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        adjacent_law = {
            **criminal_code,
            "id": 9_311,
            "title": "Закон о деятельности государственных служащих",
            "content_text": "Статья 3. Общие полномочия. Наказание определяется законом.",
            "source_url": "https://forum.majestic-rp.ru/threads/law.9311/",
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[adjacent_law, criminal_code],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(77, "Назови статью за убийство", expanded=True)

        self.assertTrue(result)
        self.assertEqual(result[0]["source_id"], criminal_code["id"])
        self.assertEqual(result[0]["reference"], "article:6.2")
        self.assertIn("6.2", result[0]["text"])
        self.assertIn("Убийство", result[0]["text"])

    async def test_thematic_legal_search_matches_inflected_offence(self) -> None:
        source = {
            "id": 9_312,
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "8.1 (F/R) Кража чужого имущества. Наказание: до 30 месяцев.\n"
                "8.2 (F/R) Грабеж с применением насилия. Наказание: до 40 месяцев."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.9312/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(77, "Какая статья за кражу?", expanded=True)

        self.assertEqual(result[0]["reference"], "article:8.1")
        self.assertIn("Кража", result[0]["text"])

    async def test_thematic_legal_search_distinguishes_giving_from_receiving_bribe(self) -> None:
        source = {
            "id": 9_313,
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "15.4 (F/R) Получение взятки должностным лицом. Наказание: до 50 месяцев.\n"
                "15.5 (F/R) Дача взятки должностному лицу. Наказание: до 40 месяцев."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.9313/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(77, "Какая статья за дачу взятки?", expanded=True)

        self.assertEqual(result[0]["reference"], "article:15.5")

    async def test_colloquial_killing_question_finds_dm_rule(self) -> None:
        source = self._project_rules_source()
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Можно ли убивать без причины по правилам сервера?",
                expanded=True,
            )

        self.assertEqual(result[0]["reference"], "clause:5.1")
        self.assertIn("DM", result[0]["text"])

    async def test_complaint_wording_finds_dm_rule_before_generic_clauses(self) -> None:
        source = self._project_rules_source()
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Игрок убил меня без причины и диалога. Составь жалобу.",
                expanded=True,
                allowed_domains=("ooc", "mixed"),
            )

        self.assertEqual(result[0]["reference"], "clause:5.1")
        self.assertIn("DM", result[0]["text"])

    async def test_search_returns_exact_article_instead_of_nearby_reference(self) -> None:
        source = {
            "id": 94,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "Глава 16.\nПреступления против правосудия.\n"
                "16.1\nПервая точная норма.\n"
                "16.1.2\nДругая вложенная норма.\n"
                "16.2\nСледующая статья."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.4/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(77, "Покажи статью 16.1 УК", expanded=True)

        self.assertEqual(result[0]["reference"], "article:16.1")
        self.assertIn("Первая точная норма", result[0]["text"])
        self.assertNotIn("Другая вложенная норма", result[0]["text"])

    async def test_exact_article_uses_the_named_law_among_duplicate_numbers(self) -> None:
        common = {
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        sources = [
            {
                **common,
                "id": 9_401,
                "title": "Закон О статусе United States Secret Service",
                "content_text": "3.1 Полномочия секретной службы.",
                "source_url": "https://forum.majestic-rp.ru/threads/usss.9401/",
            },
            {
                **common,
                "id": 9_402,
                "title": "Закон О Правительстве штата San-Andreas",
                "content_text": "3.1 Правительство формирует систему органов исполнительной власти.",
                "source_url": "https://forum.majestic-rp.ru/threads/government.9402/",
            },
            {
                **common,
                "id": 9_403,
                "title": "Закон О государственных документах штата San-Andreas",
                "content_text": "3.1 Государственный документ имеет обязательные реквизиты.",
                "source_url": "https://forum.majestic-rp.ru/threads/documents.9403/",
            },
        ]
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            # Keep the similarly named documents law first: the resolver must
            # use the title itself rather than relying on database row order.
            return_value=list(reversed(sources)),
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Что написано в статье 3.1 закона о Правительстве?",
                expanded=True,
            )

        self.assertEqual(result[0]["source_id"], 9_402)
        self.assertIn("исполнительной власти", result[0]["text"])

        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=list(reversed(sources)),
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            audit_wording = await atlas_search(
                77,
                "Покажи статью 3.1 документа Закон О Правительстве штата San-Andreas",
                expanded=True,
            )

        self.assertEqual(audit_wording[0]["source_id"], 9_402)

        similarly_named_sources = [
            {
                **common,
                "id": 9_404,
                "title": "Закон О государственных документах штата San-Andreas",
                "content_text": "3.14 Документ прекращает действие после аннулирования.",
            },
            {
                **common,
                "id": 9_405,
                "title": (
                    "Закон Об обороте оружия и государственных специальных "
                    "средств штата San-Andreas"
                ),
                "content_text": "3.14 Оружие хранится в установленном порядке.",
            },
        ]
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            # The shared ``государствен…`` stem must not make row order decide
            # which law owns the article.
            return_value=list(reversed(similarly_named_sources)),
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(
                side_effect=AtlasAIError(
                    "upstream_unavailable", "offline", retryable=True
                )
            ),
        ):
            documents_law = await atlas_search(
                77,
                "Покажи статью 3.14 документа «Закон О государственных документах штата San-Andreas»",
                expanded=True,
            )

        self.assertEqual(documents_law[0]["source_id"], 9_404)
        self.assertIn("аннулирования", documents_law[0]["text"])

    async def test_exact_article_understands_corpus_organization_abbreviation(self) -> None:
        common = {
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        sources = [
            {
                **common,
                "id": 9_411,
                "title": "Закон О Статусе San Andreas National Guard",
                "content_text": "3.1 Применение вооружённых сил.",
            },
            {
                **common,
                "id": 9_412,
                "title": "Закон О статусе United States Secret Service",
                "content_text": "3.1 Секретная служба обеспечивает охрану первых лиц штата.",
            },
        ]
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=sources,
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Что написано в статье 3.1 закона о статусе USSS?",
                expanded=True,
            )

        self.assertEqual(result[0]["source_id"], 9_412)
        self.assertIn("охрану первых лиц", result[0]["text"])

    async def test_exact_article_is_not_displaced_by_a_planner_reference(self) -> None:
        source = {
            "id": 9_314,
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "1.3 Совокупность преступлений. В тексте упоминается статья 17.3.\n"
                "17.3 Оскорбление представителя власти. Наказание: до 30 месяцев."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.9314/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Что означает статья 17.3 Уголовного кодекса?",
                expanded=True,
                query_variants=["Проверить статью 1.3 и её исключения"],
            )

        self.assertEqual(result[0]["reference"], "article:17.3")
        self.assertIn("Оскорбление представителя власти", result[0]["text"])

    async def test_exact_article_survives_forum_markup_and_nonbreaking_spaces(self) -> None:
        source = {
            "id": 95,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "# [B]Статья\u00a016.1[/B]\nТочная норма из оформленной темы.\n"
                "[SIZE=5][B]Статья 16.2[/B][/SIZE]\nСледующая норма."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.5/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(77, "Покажи статью 16.1 УК", expanded=True)

        self.assertEqual(result[0]["reference"], "article:16.1")
        self.assertIn("Точная норма из оформленной темы", result[0]["text"])
        self.assertNotIn("Следующая норма", result[0]["text"])

    async def test_arabic_chapter_request_matches_roman_forum_heading(self) -> None:
        source = {
            "id": 96,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "[CENTER][B]ГЛАВА XVI[/B][/CENTER]\nНужная глава.\n"
                "[CENTER][B]ГЛАВА XVII[/B][/CENTER]\nСледующая глава."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.6/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(77, "Напиши 16 главу УК", expanded=True)

        self.assertEqual(result[0]["reference"], "chapter:16")
        self.assertIn("Нужная глава", result[0]["text"])
        self.assertNotIn("Следующая глава", result[0]["text"])

    async def test_exact_chapter_survives_markdown_heading_decoration(self) -> None:
        source = {
            "id": 961,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный Кодекс штата San Andreas",
            "content_text": (
                "**ГЛАВА 16. ПРЕСТУПЛЕНИЯ ПРОТИВ ПРАВОСУДИЯ**\n"
                "Статья 16.1. Точная норма.\n"
                "**ГЛАВА 17. ИНЫЕ ПРЕСТУПЛЕНИЯ**\nСледующая глава."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/uk.61/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(77, "Покажи главу 16 УК", expanded=True)

        self.assertEqual(result[0]["reference"], "chapter:16")
        self.assertIn("Точная норма", result[0]["text"])
        self.assertNotIn("Следующая глава", result[0]["text"])

    async def test_lexical_fallback_matches_russian_word_forms(self) -> None:
        source = {
            "id": 962,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Процессуальный кодекс",
            "content_text": "Порядок задержания требует разъяснить гражданину основание процедуры.",
            "source_url": "https://forum.majestic-rp.ru/threads/process.62/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(77, "Меня задержали, что делать?", expanded=True)

        self.assertEqual(result[0]["source_id"], 962)
        self.assertIn("задержания", result[0]["text"])

    async def test_ic_detention_query_does_not_promote_unrelated_ooc_clauses(self) -> None:
        procedural = {
            "id": 963,
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "title": "Процессуальный Кодекс штата San Andreas",
            "content_text": (
                "Глава 4. Задержание. Сотрудник обязан назвать основание задержания и "
                "разъяснить задержанному право на защиту."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/process.63/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        event_rules = {
            "id": 964,
            "organization_id": 1,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "global",
            "federation_scope": "project",
            "title": "Правила нападения на военную базу",
            "content_text": (
                "1.1 Сотрудники могут участвовать в событии.\n"
                "1.2 Действия участников должны соответствовать правилам мероприятия."
            ),
            "source_url": "https://forum.majestic-rp.ru/threads/event.64/",
            "metadata": {"taxonomy": {"domain": "ooc", "corpus_kind": "server_rule"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[procedural, event_rules],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=lambda texts: [[0.1, 0.2] for _ in texts]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(
                77,
                "Меня задержали сотрудники LSPD, какие у меня права и что делать?",
                expanded=True,
            )

        self.assertEqual(result[0]["source_id"], 963)
        self.assertFalse(any(item.get("structured") for item in result if item["source_id"] == 964))

    async def test_exact_lookup_returns_canonical_text_without_model_rewrite(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        source = {
            "source_id": 965,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "knowledge_domain": "ic",
            "corpus_kind": "law",
            "authority_scope": "state",
            "title": "Уголовный Кодекс штата San Andreas",
            "url": "https://forum.majestic-rp.ru/threads/uk.65/",
            "text": "Глава 16. Преступления против правосудия.\n16.1 Точная норма.",
            "score": 10.0,
            "structured": True,
            "reference": "chapter:16",
            "pinpoints": ["глава 16"],
        }
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search", AsyncMock(return_value=[source])
        ), patch("modules.atlas_ai._json_request", AsyncMock()) as provider:
            result = await atlas_answer(77, "Напиши полностью главу 16 УК")

        provider.assert_not_awaited()
        self.assertIn("16.1 Точная норма", result["answer"])
        self.assertIn("[1, глава 16]", result["answer"])
        self.assertEqual(result["model_provider"], "tmod")
        self.assertEqual(result["citation_health"]["status"], "ok")

    async def test_hybrid_search_uses_saved_source_when_semantic_search_is_down(self) -> None:
        source = {
            "id": 92,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Уголовный кодекс",
            "content_text": "Уголовный кодекс определяет преступления и ответственность.",
            "source_url": "https://forum.majestic-rp.ru/threads/uk.2/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(
                side_effect=AtlasAIError(
                    "upstream_rate_limited",
                    "temporary limit",
                    retryable=True,
                )
            ),
        ):
            result = await atlas_search(77, "Что такое УК?", expanded=True)

        self.assertTrue(result)
        self.assertEqual(result[0]["source_id"], 92)

    async def test_ooc_rules_search_returns_account_transfer_clause_before_semantic_chunks(self) -> None:
        source = self._project_rules_source()
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Можно ли передавать свой аккаунт другому игроку? Укажи пункт и наказание.",
                expanded=True,
            )

        self.assertTrue(result)
        self.assertEqual(result[0]["reference"], "clause:2.2")
        self.assertIn("Запрещено передавать аккаунт", result[0]["text"])
        self.assertIn("PermBan", result[0]["text"])
        merged = _atlas_merge_source_fragments(result)
        self.assertIn("пункт 2.2", merged[0]["pinpoints"])

    async def test_ooc_rules_exact_clause_keeps_descendant_without_matching_lookalike(self) -> None:
        source = self._project_rules_source()
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(77, "Покажи пункт 2.2 правил", expanded=True)

        self.assertEqual(result[0]["reference"], "clause:2.2")
        self.assertIn("2.2.1 Вложенное пояснение", result[0]["text"])
        self.assertNotIn("2.20 Условный соседний", result[0]["text"])
        self.assertNotIn("2.3 Администрация", result[0]["text"])

    async def test_ooc_rules_search_returns_dm_definition_and_exceptions(self) -> None:
        source = self._project_rules_source()
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Что такое DM по правилам проекта и какие есть исключения из требования IC-диалога?",
                expanded=True,
            )

        self.assertEqual(result[0]["reference"], "clause:5.1")
        self.assertIn("прямое убийство", result[0]["text"])
        self.assertIn("GunBan 8 часов", result[0]["text"])
        self.assertIn("Исключение: IC диалог", result[0]["text"])

    async def test_ooc_rule_ranking_prefers_definition_over_event_exception(self) -> None:
        main_rules = self._project_rules_source()
        event_rules = {
            **self._project_rules_source(),
            "id": 9_072,
            "title": "Правила нападения на Форт-Занкудо",
            "source_url": "https://forum.majestic-rp.ru/threads/fort.9000/",
            "content_text": "1.2 На мероприятии действуют общие правила. Исключение: PG, DM.",
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[event_rules, main_rules],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(
                77,
                "Что такое DM и какое наказание предусмотрено правилами проекта?",
                expanded=True,
            )

        self.assertEqual(result[0]["source_id"], main_rules["id"])
        self.assertEqual(result[0]["reference"], "clause:5.1")

    async def test_ooc_rule_parser_splits_spaced_xenforo_numbers(self) -> None:
        source = {
            **self._project_rules_source(),
            "content_text": (
                "5.1 DM — убийство без IC причины и IC диалога. | WARN.\n"
                "5. 2 Запрещено стороннее ПО; этот пункт не относится к DM.\n"
                "5. 3 Запрещено использовать ошибки игры."
            ),
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            result = await atlas_search(77, "Что такое DM?", expanded=True)

        self.assertEqual(result[0]["reference"], "clause:5.1")
        self.assertNotIn("стороннее ПО", result[0]["text"])

    async def test_ooc_rules_search_keeps_chat_and_relatives_answers_pinpointed(self) -> None:
        source = self._project_rules_source()
        checks = (
            (
                "Какое наказание предусмотрено за прямое оскорбление родственников? Укажи пункт правил.",
                "clause:4.3",
                "HardBan 30–60 дней",
            ),
            (
                "В каком чате можно передавать OOC-информацию и что запрещено в обычном голосовом чате?",
                "clause:4.1",
                "/b, /fb, /gb и /cb",
            ),
        )
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(side_effect=AtlasAIError("upstream_unavailable", "offline", retryable=True)),
        ):
            for question, reference, expected in checks:
                result = await atlas_search(77, question, expanded=True)
                self.assertEqual(result[0]["reference"], reference)
                self.assertIn(expected, result[0]["text"])

    async def test_search_uses_all_accessible_knowledge_scopes(self) -> None:
        canonical = {
            "id": 4,
            "organization_id": 77,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "checksum": "fresh-checksum",
            "title": "Регламент",
            "source_url": None,
            "metadata": {"taxonomy": {"domain": "mixed", "corpus_kind": "procedure"}},
        }
        response = {
            "result": {
                "points": [
                    {
                        "score": 0.91,
                        "payload": {
                            "organization_id": 77,
                            "source_id": 4,
                            "project_code": "majestic-rp",
                            "federation_scope": "server",
                            "access_scope": "server:majestic-rp:phoenix-15",
                            "checksum": "fresh-checksum",
                            "title": "Регламент",
                            "text": "Текст",
                        },
                    }
                ]
            }
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[],
        ), patch("modules.atlas_ai.atlas_embed", AsyncMock(return_value=[[0.1, 0.2]])), patch(
            "modules.atlas_ai._json_request", AsyncMock(return_value=response)
        ) as request, patch(
            "modules.atlas_ai.atlas_storage.atlas_visible_knowledge_sources_by_id",
            return_value={4: canonical},
        ):
            result = await atlas_search(77, "полномочия")

        payload = request.await_args.kwargs["payload"]
        self.assertEqual(
            payload["filter"]["must"][0],
            {
                "key": "access_scope",
                "match": {
                    "any": [
                        "platform",
                        "project:majestic-rp",
                        "server:majestic-rp:phoenix-15",
                        "faction:majestic-rp:phoenix-15:lspd",
                        "workspace:majestic-rp:77:phoenix-15:lspd",
                    ]
                },
            },
        )
        self.assertEqual(result[0]["source_id"], 4)

    async def test_semantic_hit_is_rejected_when_canonical_revision_changed(self) -> None:
        canonical = {
            "id": 501,
            "organization_id": 77,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "checksum": "current-revision",
            "title": "Новая редакция",
            "source_url": "https://forum.example.org/501",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        response = {
            "result": {
                "points": [
                    {
                        "score": 0.99,
                        "payload": {
                            "source_id": 501,
                            "project_code": "majestic-rp",
                            "federation_scope": "server",
                            "access_scope": "server:majestic-rp:phoenix-15",
                            "checksum": "obsolete-revision",
                            "text": "Текст устаревшей редакции.",
                        },
                    }
                ]
            }
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[],
        ), patch(
            "modules.atlas_ai.atlas_storage.atlas_visible_knowledge_sources_by_id",
            return_value={501: canonical},
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(return_value=[[0.1, 0.2]]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value=response),
        ):
            result = await atlas_search(77, "Покажи норму")

        self.assertEqual(result, [])

    async def test_agent_domain_filter_does_not_mix_ic_material_into_ooc_complaint(self) -> None:
        source = {
            "id": 502,
            "organization_id": 77,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "federation_scope": "server",
            "title": "Уголовный кодекс",
            "content_text": "Уголовный кодекс содержит применимые составы правонарушений.",
            "source_url": "https://forum.example.org/502",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(return_value=[[0.1, 0.2]]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            result = await atlas_search(
                77,
                "Какая жалоба по правилам сервера?",
                allowed_domains=("ooc", "mixed"),
            )

        self.assertEqual(result, [])

    async def test_search_applies_server_and_faction_filters(self) -> None:
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[],
        ), patch("modules.atlas_ai.atlas_embed", AsyncMock(return_value=[[0.1, 0.2]])), patch(
            "modules.atlas_ai._json_request", AsyncMock(return_value={"result": {"points": []}})
        ) as request:
            await atlas_search(
                77,
                "порядок задержания",
                server_code="phoenix-15",
                faction_code="lspd",
            )

        scopes = request.await_args.kwargs["payload"]["filter"]["must"][0]
        self.assertEqual(scopes["key"], "access_scope")
        self.assertIn("server:majestic-rp:phoenix-15", scopes["match"]["any"])
        self.assertIn("faction:majestic-rp:phoenix-15:lspd", scopes["match"]["any"])
        self.assertIn("workspace:majestic-rp:77:phoenix-15:lspd", scopes["match"]["any"])

    async def test_search_embeds_agent_queries_independently(self) -> None:
        embedded: list[str] = []

        async def embed(texts: list[str]) -> list[list[float]]:
            embedded.extend(texts)
            return [[0.1, 0.2] for _ in texts]

        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[],
        ), patch("modules.atlas_ai.atlas_embed", side_effect=embed), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            await atlas_search(
                77,
                "основной вопрос",
                query_variants=["процесс задержания", "исключения из правила"],
            )

        self.assertEqual(
            embedded,
            ["основной вопрос", "процесс задержания", "исключения из правила"],
        )

    async def test_expanded_search_does_not_drop_late_verification_queries(self) -> None:
        embedded: list[str] = []

        async def embed(texts: list[str]) -> list[list[float]]:
            embedded.extend(texts)
            return [[0.1, 0.2] for _ in texts]

        checks = [
            "основная норма",
            "исключение",
            "компетенция",
            "срок процедуры",
            "порядок обжалования",
        ]
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[],
        ), patch("modules.atlas_ai.atlas_embed", side_effect=embed), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {"points": []}}),
        ):
            await atlas_search(
                77,
                "законность задержания",
                expanded=True,
                query_variants=checks,
            )

        self.assertEqual(embedded[:6], ["законность задержания", *checks])

    async def test_planned_queries_also_work_in_lexical_fallback(self) -> None:
        source = {
            "id": 95,
            "organization_id": 1,
            "server_code": "phoenix-15",
            "faction_code": "lspd",
            "visibility_scope": "server",
            "title": "Процессуальный кодекс",
            "content_text": "Специальное исключение разрешает прекратить процессуальное действие.",
            "source_url": "https://forum.majestic-rp.ru/threads/process.5/",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[source],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(
                side_effect=AtlasAIError(
                    "upstream_unavailable",
                    "temporary outage",
                    retryable=True,
                )
            ),
        ):
            result = await atlas_search(
                77,
                "Что делать дальше?",
                query_variants=["специальное исключение процессуальное действие"],
            )

        self.assertEqual(result[0]["source_id"], 95)
        self.assertIn("исключение", result[0]["text"])

    async def test_index_payload_contains_one_canonical_access_scope(self) -> None:
        source = {
            "id": 5,
            "organization_id": 77,
            "project_code": "majestic-rp",
            "server_code": "phoenix-15",
            "faction_code": "gov",
            "visibility_scope": "server",
            "federation_scope": "server",
            "checksum": "source-checksum",
            "title": "Общий регламент",
            "content_text": "Проверенный общий материал Phoenix длиной больше двадцати символов.",
        }
        embed = AsyncMock(return_value=[[0.1, 0.2]])
        with patch(
            "modules.atlas_ai.atlas_embed",
            embed,
        ), patch(
            "modules.atlas_ai.atlas_ensure_collection",
            AsyncMock(),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"result": {}}),
        ) as request:
            await atlas_index_source(source)

        put_call = next(call for call in request.await_args_list if call.args[0] == "PUT")
        delete_call = next(call for call in request.await_args_list if call.args[0] == "POST")
        point = put_call.kwargs["payload"]["points"][0]
        self.assertIn("Название документа: Общий регламент", embed.await_args.args[0][0])
        self.assertEqual(point["payload"]["access_scope"], "server:majestic-rp:phoenix-15")
        self.assertEqual(point["payload"]["visibility_scope"], "server")
        self.assertEqual(point["payload"]["federation_scope"], "server")
        self.assertEqual(point["payload"]["project_code"], "majestic-rp")
        self.assertEqual(point["payload"]["checksum"], "source-checksum")
        self.assertEqual(point["payload"]["knowledge_domain"], "mixed")
        self.assertEqual(point["payload"]["corpus_kind"], "procedure")
        self.assertEqual(point["payload"]["index_version"], 3)
        self.assertGreater(len(delete_call.kwargs["payload"]["points"]), 0)
        self.assertLess(
            request.await_args_list.index(put_call),
            request.await_args_list.index(delete_call),
        )

    async def test_collection_is_created_only_when_missing(self) -> None:
        missing = AtlasAIError("upstream_not_found", "missing")
        request = AsyncMock(side_effect=[missing, {"result": True}])
        with patch("modules.atlas_ai._json_request", request):
            await atlas_ensure_collection(1536)
        self.assertEqual(request.await_count, 2)
        self.assertEqual(request.await_args.kwargs["payload"]["vectors"]["size"], 1536)

        denied = AsyncMock(side_effect=AtlasAIError("upstream_error", "denied"))
        with patch("modules.atlas_ai._json_request", denied), self.assertRaises(AtlasAIError):
            await atlas_ensure_collection(1536)
        self.assertEqual(denied.await_count, 1)

    async def test_embedding_retries_temporary_rate_limit(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat/completions",
            chat_model="test/chat",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        request = AsyncMock(
            side_effect=[
                AtlasAIError("upstream_rate_limited", "limit", retryable=True),
                {"data": [{"embedding": [0.1, 0.2]}]},
            ]
        )
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai._json_request", request
        ), patch("modules.atlas_ai.asyncio.sleep", AsyncMock()) as sleep:
            vectors = await atlas_embed(["Уголовный кодекс"])

        self.assertEqual(vectors, [[0.1, 0.2]])
        self.assertEqual(request.await_count, 2)
        sleep.assert_awaited_once_with(2)

    async def test_search_treats_qdrant_gridstore_panic_as_optional_accelerator(self) -> None:
        corrupted = AtlasAIError(
            "qdrant_index_corrupted",
            "Service internal error: task panicked with OutputTooSmall",
            retryable=True,
        )
        with patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=[],
        ), patch(
            "modules.atlas_ai.atlas_embed",
            AsyncMock(return_value=[[0.1, 0.2]]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(side_effect=corrupted),
        ):
            result = await atlas_search(77, "порядок задержания")

        # The canonical database remains authoritative. A broken Qdrant
        # vector must not turn a normal question into a 5xx/retrieval refusal;
        # the caller can still answer from lexical evidence or the model's
        # grounded fallback while the index is rebuilt in the background.
        self.assertEqual(result, [])

    async def test_collection_probe_detects_missing_and_corrupt_payload(self) -> None:
        missing = AsyncMock(side_effect=AtlasAIError("upstream_not_found", "missing"))
        with patch("modules.atlas_ai._json_request", missing):
            self.assertEqual((await atlas_probe_collection())["status"], "missing")

        corrupted = AsyncMock(
            side_effect=[
                {"result": {"points_count": 1}},
                AtlasAIError(
                    "qdrant_index_corrupted",
                    "Gridstore OutputTooSmall",
                    retryable=True,
                ),
            ]
        )
        with patch("modules.atlas_ai._json_request", corrupted):
            self.assertEqual((await atlas_probe_collection())["status"], "corrupted")

        stale = AsyncMock(
            side_effect=[
                {"result": {"points_count": 12}},
                {"result": {"points": [{"payload": {"title": "Старый индекс"}}]}},
            ]
        )
        with patch("modules.atlas_ai._json_request", stale):
            self.assertEqual((await atlas_probe_collection())["status"], "stale")

        current = AsyncMock(
            side_effect=[
                {"result": {"points_count": 12}},
                {"result": {"points": [{"payload": {"index_version": 3}}]}},
            ]
        )
        with patch("modules.atlas_ai._json_request", current):
            self.assertEqual((await atlas_probe_collection())["status"], "ok")

    async def test_openrouter_answer_is_delivered_as_real_sse_deltas(self) -> None:
        requests = []
        async def completion(request: web.Request) -> web.StreamResponse:
            payload = await request.json()
            requests.append(payload)
            self.assertTrue(payload["stream"])
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(request)
            await response.write('data: {"choices":[{"delta":{"content":"Первый "}}]}\n\n'.encode())
            await response.write('data: {"choices":[{"delta":{"content":"фрагмент"}}]}\n\n'.encode())
            await response.write(b"data: [DONE]\n\n")
            await response.write_eof()
            return response

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="test/model",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        chunks: list[str] = []

        async def receive(text: str) -> None:
            chunks.append(text)

        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search",
                AsyncMock(return_value=[]),
            ):
                result = await atlas_answer_stream(77, "Ответь по частям", on_delta=receive)
                voice_result = await atlas_answer_stream(77, "Ответь по частям", on_delta=AsyncMock(), conversation_mode="voice")
        finally:
            await server.close()

        self.assertEqual(chunks, ["Первый ", "фрагмент"])
        self.assertEqual(result["answer"], "Первый фрагмент")
        self.assertEqual(result["model"], "test/model")
        self.assertEqual(result["model_provider"], "openrouter")
        self.assertEqual(voice_result["answer"], result["answer"])
        self.assertLessEqual(requests[-1]["max_tokens"], 420)
        self.assertEqual(requests[-1]["messages"][-2]["role"], "system")
        self.assertEqual(requests[-1]["messages"][-1]["role"], "user")
        self.assertIn("голосовой разговор", requests[-1]["messages"][-2]["content"])
        self.assertNotIn("голосовой разговор", requests[0]["messages"][-1]["content"])

    async def test_source_backed_stream_never_leaks_provider_retrieval_refusal(self) -> None:
        async def completion(request: web.Request) -> web.StreamResponse | web.Response:
            body = await request.json()
            refusal = "В библиотеке Atlas нет точной статьи."
            if not body.get("stream"):
                return web.json_response({"choices": [{"message": {"content": refusal}}]})
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(request)
            await response.write(
                f'data: {{"choices":[{{"delta":{{"content":"{refusal}"}}}}]}}\n\n'.encode()
            )
            await response.write(b"data: [DONE]\n\n")
            await response.write_eof()
            return response

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="test/model",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        chunks: list[str] = []

        async def receive(text: str) -> None:
            chunks.append(text)

        source = {
            "source_id": 7,
            "title": "Правила проекта",
            "url": "https://example.test/rules",
            "text": "Правила поведения участника проекта.",
            "structured": False,
            "reference": "",
            "pinpoints": [],
            "knowledge_domain": "ooc",
            "corpus_kind": "rules",
            "score": 5.0,
        }
        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search",
                AsyncMock(return_value=[source]),
            ):
                result = await atlas_answer_stream(77, "Что происходит?", on_delta=receive)
        finally:
            await server.close()

        joined = "".join(chunks)
        self.assertNotIn("библиотек", joined.casefold())
        self.assertNotIn("библиотек", result["answer"].casefold())
        self.assertIn("уточни", result["answer"].casefold())
        self.assertEqual(chunks, [result["answer"]])

    async def test_source_free_stream_does_not_flash_retrieval_refusal(self) -> None:
        active_refusal = [""]

        async def completion(request: web.Request) -> web.StreamResponse | web.Response:
            body = await request.json()
            refusal = active_refusal[0]
            if not body.get("stream"):
                return web.json_response(
                    {
                        "choices": [
                            {
                                "message": {
                                    "content": refusal
                                }
                            }
                        ]
                    }
                )
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(request)
            # Split a diagnostic like a real token stream, before its text
            # can be classified by the full refusal recognizer.
            for part in (refusal[:20], refusal[20:]):
                payload = json.dumps(
                    {"choices": [{"delta": {"content": part}}]},
                    ensure_ascii=False,
                )
                await response.write(f"data: {payload}\n\n".encode())
            await response.write(b"data: [DONE]\n\n")
            await response.write_eof()
            return response

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="test/model",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search",
                AsyncMock(return_value=[]),
            ):
                for refusal in (
                    "В библиотеке Atlas нет точной статьи по запросу.",
                    "В этом контексте нет информации по запросу.",
                    "По запросу ничего не найдено в источниках.",
                    "Атлас не знает ответа по этому вопросу.",
                ):
                    active_refusal[0] = refusal
                    chunks: list[str] = []

                    async def receive(text: str) -> None:
                        chunks.append(text)

                    result = await atlas_answer_stream(
                        77,
                        "Скажи коротко, что делать.",
                        on_delta=receive,
                    )
                    joined = "".join(chunks).casefold()
                    self.assertNotIn("библиотек", joined)
                    self.assertNotIn("контексте нет", joined)
                    self.assertNotIn("ничего не найдено", joined)
                    self.assertNotIn("атлас не знает", joined)
                    self.assertTrue(
                        any(term in result["answer"].casefold() for term in ("уточни", "опиши"))
                    )
                    self.assertEqual(chunks, [result["answer"]])
        finally:
            await server.close()

    async def test_partial_stream_continues_after_provider_token_limit(self) -> None:
        requests: list[dict] = []

        async def completion(request: web.Request) -> web.StreamResponse:
            body = await request.json()
            requests.append(body)
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(request)
            if len(requests) == 1:
                await response.write(
                    'data: {"choices":[{"delta":{"content":"Начало ответа"}}]}\n\n'.encode()
                )
                await response.write(
                    'data: {"choices":[{"delta":{},"finish_reason":"length"}]}\n\n'.encode()
                )
            else:
                await response.write(
                    'data: {"choices":[{"delta":{"content":" и завершение."}}]}\n\n'.encode()
                )
                await response.write(
                    'data: {"choices":[{"delta":{},"finish_reason":"stop"}]}\n\n'.encode()
                )
            await response.write(b"data: [DONE]\n\n")
            await response.write_eof()
            return response

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        chunks: list[str] = []

        async def receive(text: str) -> None:
            chunks.append(text)

        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search", AsyncMock(return_value=[])
            ):
                result = await atlas_answer_stream(77, "Дай полный ответ", on_delta=receive)
        finally:
            await server.close()

        self.assertEqual(result["answer"], "Начало ответа\n\n и завершение.")
        self.assertEqual(len(requests), 2)
        self.assertEqual(requests[1]["reasoning"]["effort"], "minimal")
        self.assertEqual(requests[1]["messages"][-2]["content"], "Начало ответа")
        self.assertIn("Продолжи ровно с места обрыва", requests[1]["messages"][-1]["content"])

    async def test_empty_stream_is_retried_once_as_visible_completion(self) -> None:
        requests: list[dict] = []

        async def completion(request: web.Request) -> web.StreamResponse:
            body = await request.json()
            requests.append(body)
            if body.get("stream"):
                response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
                await response.prepare(request)
                await response.write(
                    'data: {"choices":[{"delta":{"content":null,"reasoning":"скрыто"},"finish_reason":"length"}]}\n\n'.encode()
                )
                await response.write(b"data: [DONE]\n\n")
                await response.write_eof()
                return response
            return web.json_response(
                {"choices": [{"message": {"content": "Ответ после безопасного повтора"}}]}
            )

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        chunks: list[str] = []

        async def receive(text: str) -> None:
            chunks.append(text)

        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search", AsyncMock(return_value=[])
            ):
                result = await atlas_answer_stream(77, "Дай ответ", on_delta=receive)
        finally:
            await server.close()

        self.assertEqual(chunks, ["Ответ после безопасного повтора"])
        self.assertEqual(result["answer"], "Ответ после безопасного повтора")
        self.assertEqual(len(requests), 2)
        self.assertTrue(requests[0]["stream"])
        self.assertNotIn("stream", requests[1])
        self.assertEqual(requests[1]["reasoning"]["effort"], "minimal")
        self.assertGreaterEqual(requests[1]["max_tokens"], 3200)

    async def test_mid_stream_provider_error_is_not_reported_as_empty_answer(self) -> None:
        async def completion(request: web.Request) -> web.StreamResponse:
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(request)
            await response.write(
                'data: {"error":{"code":400,"message":"invalid request"},"choices":[{"delta":{"content":""},"finish_reason":"error"}]}\n\n'.encode()
            )
            await response.write_eof()
            return response

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="test/model",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )

        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search", AsyncMock(return_value=[])
            ):
                with self.assertRaises(AtlasAIError) as raised:
                    await atlas_answer_stream(77, "Дай ответ", on_delta=AsyncMock())
        finally:
            await server.close()

        self.assertEqual(raised.exception.code, "upstream_error")
        self.assertIn("invalid request", str(raised.exception))

    async def test_aristotle_stream_reports_visible_research_progress(self) -> None:
        calls: list[str] = []

        async def completion(request: web.Request) -> web.StreamResponse:
            body = await request.json()
            if not body.get("stream"):
                system = str(body["messages"][0]["content"])
                if "архитектор исследовательской команды" in system:
                    calls.append("planner")
                    return web.json_response(
                        {
                            "choices": [
                                {
                                    "message": {
                                        "content": json.dumps(
                                            {
                                                "mission": "Подготовить защиту при задержании",
                                                "agents": [
                                                    {
                                                        "agent": "Страж процедуры",
                                                        "role": "Процесс задержания",
                                                        "title": "Проверить действия сотрудника",
                                                        "task": "Сверить порядок задержания и права лица",
                                                        "search_query": "порядок задержания права задержанного",
                                                    },
                                                    {
                                                        "agent": "Контраргумент",
                                                        "role": "Риски позиции",
                                                        "title": "Проверить слабые места защиты",
                                                        "task": "Найти исключения и риски выбранной позиции",
                                                        "search_query": "исключения риски при задержании",
                                                    },
                                                ],
                                            },
                                            ensure_ascii=False,
                                        )
                                    }
                                }
                            ]
                        }
                    )
                calls.append("agent")
                return web.json_response(
                    {"choices": [{"message": {"content": "Проверка завершена [Источник 1]"}}]}
                )
            self.assertEqual(body["temperature"], 0.28)
            response = web.StreamResponse(headers={"Content-Type": "text/event-stream"})
            await response.prepare(request)
            await response.write('data: {"choices":[{"delta":{"content":"Итог [1]"}}]}\n\n'.encode())
            await response.write(b"data: [DONE]\n\n")
            await response.write_eof()
            return response

        app = web.Application()
        app.router.add_post("/chat", completion)
        server = TestServer(app)
        await server.start_server()
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url=str(server.make_url("/chat")),
            chat_model="test/model",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        progress: list[dict] = []

        async def receive(_text: str) -> None:
            return None

        async def report(event: dict) -> None:
            progress.append(event)

        sources = [{
            "source_id": 4,
            "title": "Процессуальный кодекс",
            "text": "Проверенная норма",
            "knowledge_domain": "ic",
            "corpus_kind": "law",
            "score": 0.9,
            "url": None,
        }]
        try:
            with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
                "modules.atlas_ai.atlas_search",
                AsyncMock(return_value=sources),
            ) as search:
                result = await atlas_answer_stream(
                    77,
                    "Проведи полное исследование по задержанию",
                    response_mode="aristotle",
                    on_delta=receive,
                    on_progress=report,
                )
        finally:
            await server.close()

        self.assertEqual(result["response_mode"], "aristotle")
        self.assertTrue(result["research_plan"])
        self.assertEqual(
            [step["agent"] for step in result["research_plan"]],
            ["Страж процедуры", "Контраргумент", "Аристотель"],
        )
        self.assertEqual(calls, ["planner", "agent", "agent"])
        self.assertTrue(search.await_args.kwargs["expanded"])
        self.assertIn("plan", {event["phase"] for event in progress})
        self.assertIn("evidence", {event["phase"] for event in progress})
        self.assertEqual(
            sum(event["phase"] == "agent_result" for event in progress),
            2,
        )
        self.assertEqual(progress[-1]["phase"], "complete")

    async def test_creative_answer_uses_current_history_cross_chat_memory_and_sources(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        sources = [
            {
                "source_id": 7,
                "title": "Закон",
                "url": "https://example.test/law",
                "text": "Публичная речь должна соблюдать требования закона.",
                "score": 0.91,
            }
        ]
        response = {"choices": [{"message": {"content": "Готовая убедительная речь [1]"}}]}
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search",
            AsyncMock(return_value=sources),
        ) as search, patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value=response),
        ) as request:
            result = await atlas_answer(
                77,
                "Теперь составь полную речь",
                history=[
                    {"role": "user", "content_text": "Мы обсуждаем судебную реформу"},
                    {"role": "assistant", "content_text": "Правовую основу я нашёл"},
                ],
                memory=[
                    {
                        "role": "user",
                        "thread_title": "Стиль выступления",
                        "content_text": "Предпочитаю спокойный официальный тон",
                    }
                ],
                response_mode="creative",
            )

        payload = request.await_args.kwargs["payload"]
        messages = payload["messages"]
        self.assertEqual(payload["temperature"], 0.68)
        self.assertEqual(payload["reasoning"]["effort"], "medium")
        self.assertIn("судебную реформу", search.await_args.args[1])
        self.assertTrue(any("Предпочитаю спокойный" in item["content"] for item in messages))
        self.assertTrue(any("не повторяй обращение" in item["content"] for item in messages))
        self.assertTrue(any("не переноси названия" in item["content"].casefold() for item in messages))
        self.assertTrue(any(item == {"role": "assistant", "content": "Правовую основу я нашёл"} for item in messages))
        self.assertEqual(messages[-1], {"role": "user", "content": "Теперь составь полную речь"})
        self.assertEqual(result["response_mode"], "creative")
        self.assertEqual(result["citations"][0]["source_id"], 7)

    async def test_legal_answer_uses_generated_research_contract(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        source = {
            "source_id": 8,
            "title": "Процессуальный кодекс",
            "url": "https://example.test/process",
            "text": "Обыск проводится при наличии предусмотренного кодексом основания.",
            "knowledge_domain": "ic",
            "corpus_kind": "law",
            "score": 0.94,
        }
        catalog = [{
            "id": 8,
            "title": "Процессуальный кодекс",
            "metadata": {"taxonomy": {"domain": "ic", "corpus_kind": "law"}},
        }]

        async def complete(_method, _url, *, payload, **_kwargs):
            if payload.get("response_format"):
                return {
                    "choices": [{
                        "message": {
                            "content": json.dumps(
                                {
                                    "resolved_question": "Проверить законность обыска по ордеру",
                                    "search_queries": [
                                        "Процессуальный кодекс основания обыска",
                                        "исключения и пределы ордера на обыск",
                                    ],
                                    "verification_points": [
                                        "основание обыска",
                                        "компетенция выдавшего ордер",
                                    ],
                                    "uncertainties": ["кто выдал ордер"],
                                    "answer_strategy": "Дать вывод и перечислить проверяемые условия",
                                },
                                ensure_ascii=False,
                            )
                        }
                    }]
                }
            return {"choices": [{"message": {"content": "Законность зависит от основания [1]."}}]}

        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_storage.atlas_searchable_knowledge_sources",
            return_value=catalog,
        ), patch(
            "modules.atlas_ai.atlas_search",
            AsyncMock(return_value=[source]),
        ) as search, patch(
            "modules.atlas_ai._json_request",
            side_effect=complete,
        ) as request:
            result = await atlas_answer(
                77,
                "Мне выдали ордер и провели обыск. Это было законно?",
            )

        self.assertEqual(request.await_count, 2)
        self.assertIn(
            "исключения и пределы ордера на обыск",
            search.await_args.kwargs["query_variants"],
        )
        self.assertIn(
            "Проверить законность обыска по ордеру. Проверить: основание обыска",
            search.await_args.kwargs["query_variants"],
        )
        self.assertEqual(result["intelligence"]["source"], "generated")
        self.assertEqual(result["evidence"]["matched_checks"], ["основание обыска"])
        self.assertEqual(
            result["evidence"]["open_checks"],
            ["компетенция выдавшего ордер"],
        )
        self.assertEqual(result["citation_health"]["status"], "ok")
        final_payload = request.await_args.kwargs["payload"]
        self.assertTrue(
            any("ИССЛЕДОВАТЕЛЬСКАЯ КАРТА" in item["content"] for item in final_payload["messages"])
        )
        self.assertTrue(
            any("КАРТА ДОКАЗАТЕЛЬСТВ ATLAS" in item["content"] for item in final_payload["messages"])
        )

    async def test_balanced_answer_can_help_when_search_has_no_confirmed_source(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="test/model",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        response = {"choices": [{"message": {"content": "Могу предложить творческий черновик."}}]}
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search",
            AsyncMock(return_value=[]),
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value=response),
        ) as request:
            result = await atlas_answer(77, "Помоги составить вступление к речи")

        self.assertEqual(result["answer"], "Могу предложить творческий черновик.")
        self.assertEqual(result["citations"], [])
        self.assertEqual(result["citation_health"]["status"], "no_sources")
        self.assertEqual(result["requested_response_mode"], "balanced")
        self.assertEqual(result["response_mode"], "creative")
        self.assertEqual(request.await_args.kwargs["payload"]["temperature"], 0.68)
        self.assertIn(
            "не отвечай шаблонным отказом о библиотеке",
            request.await_args.kwargs["payload"]["messages"][1]["content"],
        )

    async def test_quick_answer_has_compact_budget_and_non_repetitive_contract(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        source = {
            "source_id": 97,
            "title": "Уголовный кодекс",
            "url": None,
            "text": "Кодекс определяет преступления и ответственность.",
            "score": 0.9,
        }
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search", AsyncMock(return_value=[source])
        ), patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"choices": [{"message": {"content": "Короткий ответ [1]."}}]}),
        ) as request:
            result = await atlas_answer(77, "Что такое УК?")

        # Core glossary terms use the local bounded route: no planning or
        # provider completion can turn a two-word question into an essay.
        request.assert_not_awaited()
        self.assertTrue(result["answer"].startswith("УК — Уголовный кодекс"))
        self.assertLessEqual(len(result["answer"].split()), 30)

    async def test_complaint_prompt_forbids_invented_evidence_requirements(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        source = {
            "source_id": 97,
            "title": "Основные правила проекта",
            "url": None,
            "text": "5.1 DM — убийство без IC причины. | Ban.",
            "score": 10.0,
            "structured": True,
            "reference": "clause:5.1",
            "pinpoints": ["пункт 5.1"],
        }
        response = {"choices": [{"message": {"content": "DM — пункт 5.1. [1]"}}]}
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search", AsyncMock(return_value=[source])
        ), patch(
            "modules.atlas_ai._build_intelligence_brief", AsyncMock(return_value=None)
        ), patch("modules.atlas_ai._json_request", AsyncMock(return_value=response)) as request:
            await atlas_answer(
                77,
                "Игрок убил меня без причины. Составь жалобу.",
                model_id="atlas-complaints",
            )

        complaint_gate = "\n".join(
            str(item.get("content") or "") for item in request.await_args.kwargs["payload"]["messages"]
        )
        self.assertIn("Не утверждай отсутствие угрозы", complaint_gate)
        self.assertIn("не добавляй срок хранения доказательств", complaint_gate)

    async def test_non_stream_answer_retries_a_truncated_provider_response(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        request = AsyncMock(
            side_effect=[
                {
                    "choices": [
                        {"message": {"content": "Оборванный ответ в"}, "finish_reason": "length"}
                    ]
                },
                {
                    "choices": [
                        {"message": {"content": "Короткий завершённый ответ."}, "finish_reason": "stop"}
                    ]
                },
            ]
        )
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search", AsyncMock(return_value=[])
        ), patch("modules.atlas_ai._json_request", request):
            result = await atlas_answer(77, "Кратко ответь на вопрос")

        self.assertEqual(result["answer"], "Короткий завершённый ответ.")
        self.assertEqual(request.await_count, 2)
        retry_payload = request.await_args_list[1].kwargs["payload"]
        self.assertGreaterEqual(retry_payload["max_tokens"], 1800)
        self.assertIn("без оборванных предложений", retry_payload["messages"][-2]["content"])

    async def test_overlay_answer_skips_planners_and_uses_compact_field_contract(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        source = {
            "source_id": 97,
            "title": "Процессуальный кодекс",
            "url": None,
            "text": "Сотрудник обязан разъяснить задержанному основание задержания.",
            "score": 0.94,
        }
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search", AsyncMock(return_value=[source])
        ) as search, patch(
            "modules.atlas_ai._build_intelligence_brief", AsyncMock()
        ) as planner, patch(
            "modules.atlas_ai._generate_aristotle_plan", AsyncMock()
        ) as aristotle, patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"choices": [{"message": {"content": "Назовите основание [1]."}}]}),
        ) as request:
            result = await atlas_answer(
                77,
                "Меня задержали, что делать?",
                response_mode="aristotle",
                latency_mode="overlay",
                screen_context="data:image/png;base64,dmFsaWRhdGVk",
                user_profile={"nickname": "Saul Goodman", "rank": "Адвокат"},
            )

        planner.assert_not_awaited()
        aristotle.assert_not_awaited()
        self.assertEqual(search.await_args.kwargs["limit"], 5)
        self.assertTrue(search.await_args.kwargs["expanded"])
        payload = request.await_args.kwargs["payload"]
        self.assertLessEqual(payload["max_tokens"], 180)
        self.assertIn("Полевой интерфейс", payload["messages"][0]["content"])
        self.assertIn("18–36 слов", payload["messages"][0]["content"])
        self.assertIsInstance(payload["messages"][-1]["content"], list)
        self.assertEqual(result["latency_mode"], "overlay")
        self.assertTrue(result["screen_context_used"])
        self.assertEqual(result["depth"], "quick")

    async def test_overlay_greeting_skips_retrieval_and_screen_analysis(self) -> None:
        config = AtlasAIConfig(
            openrouter_key="test",
            openrouter_url="https://openrouter.test/chat",
            chat_model="openai/gpt-5-mini",
            embedding_model="test/embed",
            qdrant_url="http://qdrant",
            qdrant_key="",
            collection="atlas",
            referer="",
            title="Atlas",
        )
        with patch("modules.atlas_ai.atlas_ai_config", return_value=config), patch(
            "modules.atlas_ai.atlas_search", AsyncMock()
        ) as search, patch(
            "modules.atlas_ai._json_request",
            AsyncMock(return_value={"choices": [{"message": {"content": "Привет! Что на уме?"}}]}),
        ) as request:
            result = await atlas_answer(
                77,
                "Привет!",
                latency_mode="overlay",
                screen_context="data:image/png;base64,dmFsaWRhdGVk",
            )

        search.assert_not_awaited()
        request.assert_not_awaited()
        self.assertEqual(result["answer"], "Привет! Чем помочь?")
        self.assertEqual(result["model_provider"], "tmod")
        self.assertEqual(result["model"], "atlas-dialog")
        self.assertFalse(result["screen_context_used"])
        self.assertEqual(result["intent"], "social")

    def test_overlay_answer_hard_bound_prefers_complete_sentence(self) -> None:
        long_answer = (
            "Сначала остановитесь и уточните основание задержания [1]. "
            + "Это второстепенная подробность, которую оверлей не обязан озвучивать. " * 30
        )
        compact = _compact_overlay_answer(long_answer)

        self.assertLessEqual(len(compact), 461)
        self.assertLessEqual(len(compact.split()), 42)
        self.assertTrue(compact.startswith("Сначала остановитесь"))
        self.assertTrue(compact.endswith("…"))

    def test_overlay_keeps_short_canonical_clause_citation(self) -> None:
        answer = (
            "10.6\n(F/R)\nРазбойное ограбление — нападение с опасным насилием.\n"
            "Приоритет розыска 4\nНаказание: до 40 месяцев лишения свободы.\n\n"
            "[1, статья 10.6]"
        )
        compact = _compact_overlay_answer(answer)

        self.assertEqual(compact, answer)
        self.assertIn("[1, статья 10.6]", compact)
        self.assertNotIn(".…", compact)


class AtlasKnowledgeFileTests(unittest.TestCase):
    def test_plain_text_and_docx_are_extracted(self) -> None:
        plain = atlas_extract_knowledge_file(
            "Правила.md",
            "Проверенный материал для базы знаний Atlas длиной больше двадцати символов.".encode(),
        )
        document = Document()
        document.add_heading("Устав GOV", level=1)
        document.add_paragraph("Проверенный порядок работы государственного органа.")
        stream = io.BytesIO()
        document.save(stream)
        docx = atlas_extract_knowledge_file("Устав.docx", stream.getvalue())

        self.assertEqual(plain["title"], "Правила")
        self.assertIn("Устав GOV", docx["content"])
        self.assertIn("государственного органа", docx["content"])

    def test_unsupported_and_empty_files_are_rejected(self) -> None:
        with self.assertRaisesRegex(AtlasKnowledgeFileError, "Поддерживаются"):
            atlas_extract_knowledge_file("archive.zip", b"content")
        with self.assertRaisesRegex(AtlasKnowledgeFileError, "пуст"):
            atlas_extract_knowledge_file("empty.txt", b"")


class AtlasWebSurfaceTests(unittest.IsolatedAsyncioTestCase):
    async def test_overlay_crafts_returns_live_private_projection_and_etag(self) -> None:
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=SimpleNamespace(
                id=42,
                display_name="Администратор",
                guild_permissions=SimpleNamespace(administrator=True),
                roles=[],
            ),
        )

        async def authenticate(_request):
            return selected, False

        plan = {
            "id": 74,
            "stage": "crafting",
            "product_name_snapshot": "Бронепластины",
            "responsible_id": 42,
            "responsible_display": "Администратор",
            "attempts_total": 100,
            "attempts_queued": 20,
            "attempts_completed": 20,
            "product_stock": 20,
            "materials": [],
            "active_batch": None,
            "recipe": {"product_name": "Бронепластины"},
            "purchase_cost_total": 999_999,
        }
        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        with patch("modules.atlas_web.craft_storage.craft_active_plans", return_value=[plan]):
            async with TestClient(TestServer(app)) as client:
                response = await client.get("/api/atlas/overlay/crafts")
                payload = await response.json()
                cached = await client.get(
                    "/api/atlas/overlay/crafts",
                    headers={"If-None-Match": response.headers["ETag"]},
                )

        self.assertEqual(response.status, 200)
        self.assertEqual(cached.status, 304)
        self.assertEqual(payload["attention_count"], 1)
        self.assertTrue(payload["plans"][0]["mine"])
        self.assertTrue(payload["plans"][0]["needs_next_batch"])
        self.assertNotIn("purchase_cost_total", payload["plans"][0])

    async def test_overlay_transcribe_accepts_raw_webm_and_returns_compatibility_fields(self) -> None:
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=SimpleNamespace(
                id=42,
                guild_permissions=SimpleNamespace(administrator=True),
                roles=[],
            ),
        )

        async def authenticate(_request):
            return selected, False

        transcriber = SimpleNamespace(
            configured=True,
            transcribe_pcm=lambda _pcm: "Атлас, что мне делать?",
        )
        app = web.Application()
        with patch("modules.atlas_web.OpenRouterTranscriber", return_value=transcriber):
            register_atlas_web_routes(
                app,
                SimpleNamespace(get_guild=lambda guild_id: None),
                guild_id=77,
                asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
                authenticate=authenticate,
            )
        with patch(
            "modules.atlas_web._decode_overlay_audio",
            AsyncMock(return_value=b"\0" * (48_000 * 2 * 2)),
        ):
            async with TestClient(TestServer(app)) as client:
                response = await client.post(
                    "/api/atlas/overlay/transcribe",
                    data=b"browser-webm",
                    headers={
                        "Content-Type": "audio/webm;codecs=opus",
                        "X-CSRF-Token": "admin-csrf",
                    },
                )
                payload = await response.json()

        self.assertEqual(response.status, 200)
        self.assertEqual(payload["text"], "Атлас, что мне делать?")
        self.assertEqual(payload["transcript"], payload["text"])
        self.assertEqual(payload["duration_ms"], 1000)

    async def test_overlay_transcribe_rejects_oversized_raw_audio_before_decode(self) -> None:
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=SimpleNamespace(
                id=42,
                guild_permissions=SimpleNamespace(administrator=True),
                roles=[],
            ),
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application(client_max_size=8 * 1024**2)
        with patch(
            "modules.atlas_web.OpenRouterTranscriber",
            return_value=SimpleNamespace(configured=True),
        ):
            register_atlas_web_routes(
                app,
                SimpleNamespace(get_guild=lambda guild_id: None),
                guild_id=77,
                asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
                authenticate=authenticate,
            )
        with patch(
            "modules.atlas_web._decode_overlay_audio",
            AsyncMock(),
        ) as decode:
            async with TestClient(TestServer(app)) as client:
                response = await client.post(
                    "/api/atlas/overlay/transcribe",
                    data=io.BytesIO(b"x" * (6 * 1024**2 + 1)),
                    headers={
                        "Content-Type": "audio/webm",
                        "X-CSRF-Token": "admin-csrf",
                    },
                )
                payload = await response.json()

        self.assertEqual(response.status, 400)
        self.assertEqual(payload["error"], "atlas_overlay_audio_size_invalid")
        decode.assert_not_awaited()

    async def test_overlay_context_is_csrf_protected_and_stream_uses_owned_character_scope(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-overlay-web-test.db"
        storage.init_db()
        atlas_repository.atlas_upsert_project(42, code="project-b", name="Project B")
        atlas_repository.atlas_upsert_server(
            42,
            code="project-b-15",
            name="Phoenix",
            number=15,
            project_code="project-b",
        )
        character = storage.add_profile_character(77, 42, "Saul Goodman", "263345")
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=SimpleNamespace(
                id=42,
                display_name="Администратор",
                guild_permissions=SimpleNamespace(administrator=True),
                roles=[],
            ),
        )

        async def authenticate(_request):
            return selected, False

        async def stream_answer(_organization_id, _question, *, on_delta, **kwargs):
            await on_delta("Полевой ответ [1].")
            stream_answer.organization_id = _organization_id
            stream_answer.kwargs = kwargs
            return {
                "answer": "Полевой ответ [1].",
                "citations": [],
                "model": "atlas-tvr-a",
                "model_provider": "openrouter",
                "model_release": "base",
                "project_code": "project-b",
                "server_code": "project-b-15",
                "faction_code": "fib",
                "response_mode": "balanced",
                "requested_response_mode": "balanced",
                "latency_mode": "overlay",
                "latency_ms": 8,
            }

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        try:
            with patch("modules.atlas_web.atlas_answer_stream", stream_answer):
                async with TestClient(TestServer(app)) as client:
                    rejected = await client.post(
                        "/api/atlas/overlay/context",
                        json={
                            "character_id": character.id,
                            "server_code": "project-b-15",
                            "faction_code": "fib",
                        },
                    )
                    voice_rejected = await client.post(
                        "/api/atlas/overlay/transcribe",
                        data=FormData(),
                    )
                    configured = await client.post(
                        "/api/atlas/overlay/context",
                        json={
                            "character_id": character.id,
                            "server_code": "project-b-15",
                            "faction_code": "fib",
                            "rank": "Special Agent",
                            "screen_context_enabled": True,
                        },
                        headers={"X-CSRF-Token": "admin-csrf"},
                    )
                    context = await client.get("/api/atlas/overlay/context")
                    streamed = await client.post(
                        "/api/atlas/chat/stream",
                        json={
                            "question": "Что делать при задержании?",
                            "latency_mode": "overlay",
                            "character_id": character.id,
                            "faction_code": "lspd",
                            "screen_context": "data:image/png;base64,"
                            + base64.b64encode(b"\x89PNG\r\n\x1a\nframe").decode(),
                        },
                        headers={
                            "X-CSRF-Token": "admin-csrf",
                            "X-Idempotency-Key": "overlay-stream-1",
                        },
                    )
                    context_payload = await context.json()
                    await streamed.text()

            self.assertEqual(rejected.status, 403)
            self.assertEqual(voice_rejected.status, 403)
            self.assertEqual(configured.status, 200)
            self.assertEqual(context_payload["selected_character"]["faction_code"], "fib")
            self.assertNotIn("csrf_token", context_payload)
            self.assertEqual(streamed.status, 200)
            project_dashboard = atlas_repository.atlas_dashboard(
                77,
                42,
                "Администратор",
                project_code="project-b",
            )
            self.assertEqual(
                stream_answer.organization_id,
                int(project_dashboard["organization"]["id"]),
            )
            self.assertEqual(stream_answer.kwargs["server_code"], "project-b-15")
            self.assertEqual(stream_answer.kwargs["faction_code"], "fib")
            self.assertEqual(stream_answer.kwargs["latency_mode"], "overlay")
            self.assertEqual(stream_answer.kwargs["user_profile"]["nickname"], "Saul Goodman")
            self.assertTrue(stream_answer.kwargs["screen_context"].startswith("data:image/png"))
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()

    async def test_empty_chat_question_is_rejected_before_ai_or_stream_start(self) -> None:
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=SimpleNamespace(
                id=42,
                guild_permissions=SimpleNamespace(administrator=True),
                roles=[],
            ),
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        headers = {"X-CSRF-Token": "admin-csrf"}
        with patch("modules.atlas_web.atlas_answer", AsyncMock()) as answer, patch(
            "modules.atlas_web.atlas_answer_stream", AsyncMock()
        ) as stream:
            async with TestClient(TestServer(app)) as client:
                regular = await client.post(
                    "/api/atlas/chat",
                    json={"question": "   "},
                    headers={**headers, "X-Idempotency-Key": "empty-regular"},
                )
                streamed = await client.post(
                    "/api/atlas/chat/stream",
                    json={"question": ""},
                    headers={**headers, "X-Idempotency-Key": "empty-stream"},
                )
                regular_payload = await regular.json()
                streamed_payload = await streamed.json()

        self.assertEqual(regular.status, 400)
        self.assertEqual(streamed.status, 400)
        self.assertEqual(regular_payload["error"], "question_required")
        self.assertEqual(streamed_payload["error"], "question_required")
        answer.assert_not_awaited()
        stream.assert_not_awaited()

    async def test_atlas_surface_is_registered_and_api_requires_login(self) -> None:
        bot = SimpleNamespace(get_guild=lambda guild_id: None)
        app = create_consensus_web_app(bot, guild_id=77)
        async with TestClient(TestServer(app)) as client:
            page = await client.get("/atlas")
            bootstrap = await client.get("/api/atlas/bootstrap")
            desktop_page = await client.get(
                "/atlas",
                headers={"User-Agent": "T-Mod QA TModDesktop/0.3.5"},
            )
            public_api = await client.get(
                "/api/atlas/bootstrap",
                headers={"Host": "dash.tvr.lat"},
            )
            desktop_api = await client.get(
                "/api/atlas/bootstrap",
                headers={
                    "Host": "dash.tvr.lat",
                    "User-Agent": "T-Mod QA TModDesktop/0.3.5",
                },
            )

            self.assertEqual(page.status, 200)
            self.assertIn("T-Mod Atlas", await page.text())
            self.assertIn("Только T-Mod Desktop", await page.text())
            self.assertIn('id="atlas-app"', await desktop_page.text())
            self.assertEqual(bootstrap.status, 401)
            self.assertIn((await bootstrap.json())["error"], {"unauthorized", "atlas_login_required"})
            self.assertEqual(public_api.status, 403)
            self.assertEqual((await public_api.json())["error"], "atlas_desktop_required")
            self.assertEqual(desktop_api.status, 401)

    async def test_non_admin_receives_only_closed_preview(self) -> None:
        member = SimpleNamespace(
            id=42,
            display_name="Участник",
            guild_permissions=SimpleNamespace(administrator=False),
            roles=[],
        )
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Участник",
            csrf_token="preview-csrf",
            member=member,
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        async with TestClient(TestServer(app)) as client:
            response = await client.get("/api/atlas/bootstrap")
            payload = await response.json()
            forbidden = await client.post(
                "/api/atlas/documents",
                json={"title": "Скрытый документ"},
                headers={"X-CSRF-Token": "preview-csrf"},
            )

        self.assertTrue(payload["preview"])
        self.assertEqual(payload["catalog"]["servers"][0]["label"], "Phoenix (15)")
        self.assertEqual(
            {item["code"] for item in payload["catalog"]["factions"]},
            {"lspd", "lscsd", "fib", "gov", "sang", "ems", "wn"},
        )
        self.assertNotIn("documents", payload)
        self.assertEqual(forbidden.status, 403)

    async def test_admin_can_create_and_read_continuity_timeline(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-timeline-web-test.db"
        storage.init_db()
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=SimpleNamespace(
                id=42,
                display_name="Администратор",
                guild_permissions=SimpleNamespace(administrator=True),
                roles=[],
            ),
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        try:
            async with TestClient(TestServer(app)) as client:
                headers = {
                    "X-CSRF-Token": "admin-csrf",
                    "X-Idempotency-Key": "timeline-event-1",
                }
                created = await client.post(
                    "/api/atlas/timeline",
                    json={
                        "title": "Ситуация у штаба",
                        "event_kind": "incident",
                        "importance": "critical",
                        "summary": "Контекст сохранён для дальнейшего дела.",
                    },
                    headers=headers,
                )
                repeated = await client.post(
                    "/api/atlas/timeline",
                    json={"title": "Не должно дублироваться"},
                    headers=headers,
                )
                listed = await client.get("/api/atlas/timeline")
                payload = await listed.json()

                patched = await client.patch(
                    f"/api/atlas/timeline/{payload['items'][0]['id']}",
                    json={"status": "resolved", "expected_version": 1},
                    headers={"X-CSRF-Token": "admin-csrf"},
                )
                patched_payload = await patched.json()

            self.assertEqual(created.status, 201)
            self.assertEqual(repeated.status, 201)
            self.assertEqual(payload["summary"]["attention"], 1)
            self.assertEqual(len(payload["items"]), 1)
            self.assertEqual(payload["items"][0]["title"], "Ситуация у штаба")
            self.assertIsNone(payload["next_cursor"])
            self.assertEqual(patched.status, 200)
            self.assertEqual(patched_payload["event"]["status"], "resolved")
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()

    async def test_admin_can_upload_scoped_knowledge_file(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-web-test.db"
        storage.init_db()
        member = SimpleNamespace(
            id=42,
            display_name="Администратор",
            guild_permissions=SimpleNamespace(administrator=True),
            roles=[],
        )
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=member,
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application(client_max_size=10 * 1024 * 1024)
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        form = FormData()
        form.add_field("server_code", "phoenix-15")
        form.add_field("faction_code", "gov")
        form.add_field("visibility_scope", "global")
        form.add_field("source_kind", "regulation")
        form.add_field(
            "file",
            "Проверенный регламент Government для Atlas AI длиной больше двадцати символов.".encode(),
            filename="Регламент GOV.txt",
            content_type="text/plain",
        )
        try:
            with patch("modules.atlas_web.atlas_index_source", AsyncMock(return_value=["point-1"])):
                async with TestClient(TestServer(app)) as client:
                    uploaded = await client.post(
                        "/api/atlas/knowledge/upload",
                        data=form,
                        headers={"X-CSRF-Token": "admin-csrf", "X-Idempotency-Key": "upload-1"},
                    )
                    self.assertEqual(uploaded.status, 202, await uploaded.text())
                    await asyncio.sleep(0.05)
                    listed = await client.get(
                        "/api/atlas/knowledge?server_code=phoenix-15&faction_code=gov"
                    )
                    payload = await listed.json()
            self.assertEqual(payload["items"][0]["faction_code"], "gov")
            self.assertEqual(payload["items"][0]["visibility_scope"], "global")
            self.assertEqual(payload["items"][0]["federation_scope"], "project")
            self.assertEqual(payload["items"][0]["original_filename"], "Регламент GOV.txt")
            self.assertEqual(payload["items"][0]["status"], "indexed")
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()

    async def test_admin_can_import_and_auto_classify_authenticated_forum_thread(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-forum-import-test.db"
        storage.init_db()
        member = SimpleNamespace(
            id=42,
            display_name="Администратор",
            guild_permissions=SimpleNamespace(administrator=True),
            roles=[],
        )
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=member,
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        snapshot = AtlasForumSnapshot(
            url="https://forum.majestic-rp.ru/threads/ustav-gov.700/",
            title="Устав GOV",
            content="Настоящий устав определяет полномочия и порядок службы Government.",
            author="Author",
        )
        listing_snapshot = AtlasForumSnapshot(
            url="https://forum.majestic-rp.ru/threads/general-rule.701/",
            title="Общие правила сервера",
            content="Полная редакция общих правил сервера для всех организаций Phoenix.",
            author="Forum Admin",
        )
        try:
            with patch(
                "modules.atlas_forum_sync.AtlasForumSyncRunner.fetch_thread",
                AsyncMock(return_value=snapshot),
            ) as fetch_thread, patch(
                "modules.atlas_forum_sync.AtlasForumSyncRunner.trigger",
                return_value=True,
            ) as trigger, patch(
                "modules.atlas_forum_sync.AtlasForumSyncRunner.fetch_listing",
                AsyncMock(
                    return_value=AtlasForumScrapeBatch((listing_snapshot,), True)
                ),
            ) as fetch_listing, patch(
                "modules.atlas_web.atlas_index_source",
                AsyncMock(return_value=["point-1"]),
            ):
                async with TestClient(TestServer(app)) as client:
                    response = await client.post(
                        "/api/atlas/knowledge/import-forum",
                        json={
                            "source_url": snapshot.url,
                            "server_code": "phoenix-15",
                            "faction_code": "gov",
                            "visibility_scope": "faction",
                        },
                        headers={
                            "X-CSRF-Token": "admin-csrf",
                            "X-Idempotency-Key": "forum-import-1",
                        },
                    )
                    payload = await response.json()
                    bulk_response = await client.post(
                        "/api/atlas/knowledge/import-forum",
                        json={
                            "source_url": (
                                "https://forum.majestic-rp.ru/forums/"
                                "zakonodatel-naya-baza.1213/"
                            )
                        },
                        headers={
                            "X-CSRF-Token": "admin-csrf",
                            "X-Idempotency-Key": "forum-import-bulk-1",
                        },
                    )
                    bulk_payload = await bulk_response.json()
                    general_response = await client.post(
                        "/api/atlas/knowledge/import-forum",
                        json={
                            "source_url": (
                                "https://forum.majestic-rp.ru/forums/"
                                "general-server-rules/"
                            ),
                            "server_code": "phoenix-15",
                            "faction_code": "gov",
                            "visibility_scope": "global",
                            "knowledge_domain": "ooc",
                            "corpus_kind": "server_rule",
                        },
                        headers={
                            "X-CSRF-Token": "admin-csrf",
                            "X-Idempotency-Key": "forum-import-general-1",
                        },
                    )
                    general_payload = await general_response.json()
                    general_job = None
                    for _ in range(100):
                        general_job = atlas_job_repository.atlas_job_get(
                            int(general_payload["job"]["id"])
                        )
                        if general_job and general_job["status"] in {
                            "succeeded",
                            "failed",
                            "cancelled",
                        }:
                            break
                        await asyncio.sleep(0.05)

            self.assertEqual(response.status, 202, payload)
            self.assertTrue(payload["queued"])
            self.assertEqual(payload["job"]["job_type"], "atlas.forum.thread.v1")
            fetch_thread.assert_awaited_once_with(snapshot.url)
            self.assertEqual(bulk_response.status, 202, bulk_payload)
            self.assertTrue(bulk_payload["bulk"])
            trigger.assert_called_once_with()
            self.assertEqual(general_response.status, 202, general_payload)
            self.assertTrue(general_payload["bulk"])
            self.assertIn("каждую тему", general_payload["message"])
            self.assertIsNotNone(general_job)
            self.assertEqual(general_job["status"], "succeeded", general_job)
            fetch_listing.assert_awaited_once_with(
                "https://forum.majestic-rp.ru/forums/general-server-rules/"
            )
            sources = atlas_repository.atlas_searchable_knowledge_sources(
                int(payload["job"]["organization_id"]),
                server_code="phoenix-15",
                faction_code="gov",
            )
            by_title = {item["title"]: item for item in sources}
            self.assertIn("Общие правила сервера", by_title)
            self.assertIn("Устав GOV", by_title)
            self.assertEqual(by_title["Устав GOV"]["faction_code"], "gov")
            self.assertEqual(by_title["Устав GOV"]["metadata"]["taxonomy"]["domain"], "ic")
            self.assertEqual(
                by_title["Устав GOV"]["metadata"]["taxonomy"]["corpus_kind"],
                "charter",
            )
            rules = by_title["Общие правила сервера"]
            self.assertEqual(rules["federation_scope"], "project")
            self.assertEqual(rules["metadata"]["taxonomy"]["domain"], "ooc")
            self.assertEqual(rules["metadata"]["taxonomy"]["corpus_kind"], "server_rule")
            feed = atlas_repository.atlas_forum_sync_status(77)
            self.assertIsNotNone(feed)
            self.assertTrue(str(feed["feed_key"]).startswith("majestic-rp:manual-"))
            self.assertEqual(feed["knowledge_domain"], "ooc")
            self.assertEqual(feed["corpus_kind"], "server_rule")
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()

    async def test_chat_endpoint_continues_thread_and_exposes_owned_history(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-chat-test.db"
        storage.init_db()
        member = SimpleNamespace(
            id=42,
            display_name="Администратор",
            guild_permissions=SimpleNamespace(administrator=True),
            roles=[],
        )
        selected = ConsensusWebPrincipal(
            user_id=42,
            guild_id=77,
            display_name="Администратор",
            csrf_token="admin-csrf",
            member=member,
        )

        async def authenticate(_request):
            return selected, False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        answer = AsyncMock(
            return_value={
                "answer": "Продолжение ответа",
                "citations": [],
                "model": "test/model",
                "model_provider": "openrouter",
                "model_release": "base",
                "project_code": "majestic-rp",
                "server_code": "phoenix-15",
                "faction_code": "lspd",
                "response_mode": "creative",
                "latency_ms": 12,
            }
        )
        async def stream_answer(_organization_id, _question, *, on_delta, **_kwargs):
            await on_delta("Поток")
            await on_delta("овый ответ")
            return {
                "answer": "Потоковый ответ",
                "citations": [],
                "model": "atlas-tvr-a",
                "model_provider": "openrouter",
                "model_release": "base",
                "project_code": "majestic-rp",
                "server_code": "phoenix-15",
                "faction_code": "lspd",
                "response_mode": "balanced",
                "requested_response_mode": "balanced",
                "latency_ms": 8,
            }

        headers = {"X-CSRF-Token": "admin-csrf"}
        try:
            with patch("modules.atlas_web.atlas_answer", answer), patch(
                "modules.atlas_web.atlas_answer_stream",
                stream_answer,
            ):
                async with TestClient(TestServer(app)) as client:
                    first = await client.post(
                        "/api/atlas/chat",
                        json={"question": "Подготовь речь", "response_mode": "creative"},
                        headers={**headers, "X-Idempotency-Key": "chat-1"},
                    )
                    first_payload = await first.json()
                    second = await client.post(
                        "/api/atlas/chat",
                        json={
                            "question": "Сделай её короче",
                            "thread_id": first_payload["thread_id"],
                            "response_mode": "creative",
                        },
                        headers={**headers, "X-Idempotency-Key": "chat-2"},
                    )
                    streamed = await client.post(
                        "/api/atlas/chat/stream",
                        json={
                            "question": "Покажи поток",
                            "thread_id": first_payload["thread_id"],
                        },
                        headers={**headers, "X-Idempotency-Key": "chat-3"},
                    )
                    stream_events = [
                        json.loads(line.removeprefix("data:"))
                        for line in (await streamed.text()).splitlines()
                        if line.strip().startswith(("data:", "{"))
                    ]
                    threads = await client.get("/api/atlas/threads")
                    detail = await client.get(
                        f"/api/atlas/threads/{first_payload['thread_id']}"
                    )
                    threads_payload = await threads.json()
                    detail_payload = await detail.json()

            self.assertEqual(first.status, 200)
            self.assertEqual(second.status, 200)
            self.assertEqual(streamed.status, 200)
            self.assertEqual(
                "".join(item.get("text", "") for item in stream_events if item["type"] == "delta"),
                "Потоковый ответ",
            )
            self.assertEqual(stream_events[-1]["type"], "done")
            self.assertEqual(len(threads_payload["items"]), 1)
            self.assertEqual(len(detail_payload["messages"]), 6)
            second_history = answer.await_args_list[1].kwargs["history"]
            self.assertEqual(
                [item["role"] for item in second_history],
                ["user", "assistant"],
            )
            self.assertEqual(answer.await_args_list[1].kwargs["response_mode"], "creative")
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()

    async def test_startup_rebuilds_corrupt_search_index_from_saved_sources(self) -> None:
        old_data_dir = storage.DATA_DIR
        old_database_file = storage.DATABASE_FILE
        temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        storage.DATA_DIR = Path(temp_dir.name)
        storage.DATABASE_FILE = storage.DATA_DIR / "atlas-rebuild-test.db"
        storage.init_db()
        dashboard = atlas_repository.atlas_dashboard(77, 42, "Администратор")
        source = atlas_repository.atlas_add_knowledge(
            int(dashboard["organization"]["id"]),
            42,
            title="Сохранённый регламент",
            content="Проверенный регламент остаётся в Atlas и восстанавливает поисковый индекс.",
        )

        async def authenticate(_request):
            return None, False

        app = web.Application()
        register_atlas_web_routes(
            app,
            SimpleNamespace(get_guild=lambda guild_id: None),
            guild_id=77,
            asset_dir=Path(__file__).resolve().parents[1] / "web" / "atlas",
            authenticate=authenticate,
        )
        probe = AsyncMock(return_value={"status": "corrupted", "points_count": None})
        reset = AsyncMock(return_value=None)
        index = AsyncMock(return_value=["restored-point"])
        try:
            with patch("modules.atlas_web.atlas_probe_collection", probe), patch(
                "modules.atlas_web.atlas_reset_collection", reset
            ), patch("modules.atlas_web.atlas_index_source", index):
                async with TestClient(TestServer(app)):
                    for _ in range(50):
                        if index.await_count:
                            break
                        await asyncio.sleep(0.02)

            reset.assert_awaited_once()
            index.assert_awaited_once()
            self.assertEqual(index.await_args.args[0]["id"], source["id"])
            rebuilt = atlas_repository.atlas_knowledge_sources(
                int(dashboard["organization"]["id"])
            )[0]
            self.assertEqual(rebuilt["status"], "indexed")
        finally:
            storage.DATA_DIR = old_data_dir
            storage.DATABASE_FILE = old_database_file
            temp_dir.cleanup()


if __name__ == "__main__":
    unittest.main()
